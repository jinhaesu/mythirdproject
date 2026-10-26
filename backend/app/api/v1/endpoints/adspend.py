"""광고비 일보 API — 매체 마스터 + 일별 기입 그리드 + 월 예산(Limit/사업계획).

구글시트 "광고비 일보"의 시스템화. 금액은 원 단위(VAT 포함).
auto_source가 있는 매체(meta/naver_sa)는 일별 금액을 수집 스냅샷에서
자동으로 채우고 수동 기입을 막는다.
"""
import asyncio
import json
import logging
import os
import re
from datetime import date, datetime, timedelta
from typing import Optional

import httpx
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import delete, func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.v1.endpoints.auth import get_current_user
from app.db.database import get_db
from app.models import (
    AdMedia,
    AdMediaAlias,
    AdMediaBudget,
    AdMediaGroup,
    AdMediaSpendDaily,
    ChannelSpendDaily,
    MetaInsightDaily,
)
from app.models.user import User
from app.services import adspend_xlsx as xl

logger = logging.getLogger(__name__)
router = APIRouter()

AUTO_SOURCES = {"meta", "naver_sa"}


def _month_add(month: str, delta: int) -> str:
    y, m = int(month[:4]), int(month[5:7])
    total = y * 12 + (m - 1) + delta
    return f"{total // 12:04d}-{total % 12 + 1:02d}"


def _month_range(month: str):
    y, m = int(month[:4]), int(month[5:7])
    start = date(y, m, 1)
    end = date(y + (m == 12), m % 12 + 1, 1)
    return start, end


async def _auto_daily(db: AsyncSession, source: str, d_from: date, d_to: date) -> dict:
    """자동 수집 소스의 일별 금액 {date: amount(원)}."""
    if source == "meta":
        rows = (await db.execute(
            select(MetaInsightDaily.date, func.coalesce(func.sum(MetaInsightDaily.spend), 0))
            .where(MetaInsightDaily.level == "campaign",
                   MetaInsightDaily.date >= d_from, MetaInsightDaily.date < d_to)
            .group_by(MetaInsightDaily.date)
        )).all()
    elif source == "naver_sa":
        rows = (await db.execute(
            select(ChannelSpendDaily.date, func.coalesce(func.sum(ChannelSpendDaily.spend), 0))
            .where(ChannelSpendDaily.channel == "naver_sa",
                   ChannelSpendDaily.date >= d_from, ChannelSpendDaily.date < d_to)
            .group_by(ChannelSpendDaily.date)
        )).all()
    else:
        return {}
    return {d.isoformat(): float(s) for d, s in rows if float(s)}


# ─── 그룹(판매채널) 순서·색 ──────────────────────────────────────────────────
# 보드 그룹 = ad_media.inflow. 순서/색은 ad_media_groups(시트 포맷 그대로)에서,
# 그룹 내 매체 순서는 ad_media.sort_order. 설정 없는 inflow는 맨 뒤(이름순).

async def _group_meta(db: AsyncSession) -> dict:
    """{inflow: AdMediaGroup}"""
    return {g.name: g for g in (await db.execute(select(AdMediaGroup))).scalars().all()}


def _media_sort_key(groups: dict):
    def key(m):
        g = groups.get(m.inflow) if m.inflow else None
        if g is not None:
            gk = (0, g.sort_order, "")
        elif m.inflow:
            gk = (1, 0, m.inflow)
        else:
            gk = (2, 0, "")  # 유입 미지정(기타)은 맨 끝
        return (*gk, m.sort_order or 0, m.id)
    return key


def _groups_payload(groups: dict, media: list) -> list:
    """보드/매체추가용 그룹 목록 — 설정된 그룹(시트 순서) + 설정 없는 활성 inflow(이름순)."""
    out = [{"name": g.name, "sort_order": g.sort_order, "color": g.color, "owner": g.owner}
           for g in sorted(groups.values(), key=lambda g: (g.sort_order, g.name))]
    extra = sorted({m.inflow for m in media if m.inflow and m.inflow not in groups})
    out += [{"name": n, "sort_order": None, "color": None, "owner": None} for n in extra]
    return out


class GroupIn(BaseModel):
    name: str = Field(..., max_length=100)
    sort_order: int = 0
    color: Optional[str] = Field(None, max_length=9)
    owner: Optional[str] = Field(None, max_length=50)


@router.get("/groups")
async def list_groups(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    groups = await _group_meta(db)
    media = (await db.execute(select(AdMedia).where(AdMedia.active.is_(True)))).scalars().all()
    return _groups_payload(groups, media)


@router.put("/groups")
async def upsert_groups(
    payload: list[GroupIn],
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """그룹 순서·색 일괄 설정 (이름 기준 upsert). 목록에 없는 기존 그룹은 그대로 둔다."""
    import re as _re
    groups = await _group_meta(db)
    for g in payload:
        name = _norm_name(g.name)
        if not name:
            continue
        if g.color and not _re.fullmatch(r"#[0-9A-Fa-f]{6}", g.color):
            raise HTTPException(status_code=422, detail=f"color는 #RRGGBB 형식: {g.color}")
        row = groups.get(name)
        if not row:
            row = AdMediaGroup(name=name)
            db.add(row)
            groups[name] = row
        row.sort_order = g.sort_order
        row.color = g.color.upper() if g.color else None
        row.owner = g.owner
    await db.commit()
    return {"count": len(payload)}


class MediaOrderItem(BaseModel):
    id: int
    sort_order: int


@router.put("/media-order")
async def set_media_order(
    payload: list[MediaOrderItem],
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """매체 sort_order(그룹 내 순서)만 일괄 변경 — 다른 필드는 건드리지 않는다."""
    ids = [p.id for p in payload]
    media = {m.id: m for m in (await db.execute(select(AdMedia).where(AdMedia.id.in_(ids)))).scalars().all()}
    missing = [i for i in ids if i not in media]
    if missing:
        raise HTTPException(status_code=404, detail=f"없는 매체 id: {missing[:10]}")
    for p in payload:
        media[p.id].sort_order = p.sort_order
    await db.commit()
    return {"updated": len(payload)}


# ─── 매체 마스터 ─────────────────────────────────────────────────────────────

class MediaIn(BaseModel):
    name: str = Field(..., max_length=200)
    group_name: Optional[str] = Field(None, max_length=100)
    inflow: Optional[str] = Field(None, max_length=100)
    auto_source: Optional[str] = Field(None, max_length=20)  # meta|naver_sa
    sort_order: int = 0
    memo: Optional[str] = Field(None, max_length=300)


@router.get("/media")
async def list_media(
    include_inactive: bool = False,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    q = select(AdMedia)
    if not include_inactive:
        q = q.where(AdMedia.active.is_(True))
    rows = (await db.execute(q)).scalars().all()
    rows = sorted(rows, key=_media_sort_key(await _group_meta(db)))
    return [{
        "id": m.id, "name": m.name, "group_name": m.group_name, "inflow": m.inflow,
        "auto_source": m.auto_source, "sort_order": m.sort_order,
        "active": m.active, "memo": m.memo,
    } for m in rows]


@router.post("/media")
async def create_media(
    data: MediaIn,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    if data.auto_source and data.auto_source not in AUTO_SOURCES:
        raise HTTPException(status_code=422, detail=f"auto_source는 {sorted(AUTO_SOURCES)} 중 하나")
    fields = data.model_dump()
    fields["name"] = _norm_name(fields["name"])
    fields["inflow"] = _norm_name(fields.get("inflow") or "") or None
    fields["group_name"] = _norm_name(fields.get("group_name") or "") or None
    if not fields["name"]:
        raise HTTPException(status_code=422, detail="매체명을 입력해 주세요")
    if not fields.get("sort_order"):
        # 그룹(inflow) 맨 끝에 붙인다
        q = select(func.max(AdMedia.sort_order)).where(AdMedia.active.is_(True))
        q = q.where(AdMedia.inflow == fields["inflow"]) if fields["inflow"] else q.where(AdMedia.inflow.is_(None))
        fields["sort_order"] = int((await db.execute(q)).scalar() or 0) + 10
    if fields["inflow"]:
        # 새 그룹이면 그룹 목록 맨 끝에 등록 (이후에도 그 자리 유지)
        exists = (await db.execute(
            select(AdMediaGroup.id).where(AdMediaGroup.name == fields["inflow"])
        )).scalar_one_or_none()
        if exists is None:
            gmax = (await db.execute(select(func.max(AdMediaGroup.sort_order)))).scalar() or 0
            db.add(AdMediaGroup(name=fields["inflow"], sort_order=int(gmax) + 10))
    m = AdMedia(**fields)
    db.add(m)
    await db.commit()
    await db.refresh(m)
    return {"id": m.id}


@router.delete("/media/{media_id}")
async def delete_media(
    media_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """매체(거래처) 제거 — 기입/예산 이력이 없으면 완전 삭제, 있으면 비활성화(이력 보존).

    비활성 매체는 보드·집계에서 빠지고, 과거 월 데이터는 DB에 남는다.
    """
    m = (await db.execute(select(AdMedia).where(AdMedia.id == media_id))).scalar_one_or_none()
    if not m:
        raise HTTPException(status_code=404, detail="매체를 찾을 수 없습니다")
    has_entries = (await db.execute(
        select(func.count(AdMediaSpendDaily.id)).where(AdMediaSpendDaily.media_id == media_id)
    )).scalar() or 0
    has_budgets = (await db.execute(
        select(func.count(AdMediaBudget.id)).where(AdMediaBudget.media_id == media_id)
    )).scalar() or 0
    if has_entries or has_budgets:
        m.active = False
        await db.commit()
        return {"id": media_id, "deleted": False, "deactivated": True,
                "detail": f"기입 {has_entries}건·예산 {has_budgets}건 이력이 있어 비활성화했습니다 (데이터 보존)"}
    await db.delete(m)
    await db.commit()
    return {"id": media_id, "deleted": True, "deactivated": False}


@router.patch("/media/{media_id}")
async def update_media(
    media_id: int,
    data: MediaIn,
    active: Optional[bool] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    m = (await db.execute(select(AdMedia).where(AdMedia.id == media_id))).scalar_one_or_none()
    if not m:
        raise HTTPException(status_code=404, detail="매체를 찾을 수 없습니다")
    for k, v in data.model_dump().items():
        setattr(m, k, v)
    if active is not None:
        m.active = active
    await db.commit()
    return {"id": m.id, "active": m.active}


# ─── 일보 보드 (월 그리드) ───────────────────────────────────────────────────

@router.get("/board")
async def spend_board(
    month: str = Query(..., min_length=7, max_length=7),
    scope: str = Query("month", pattern="^(month|all)$"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """해당 월의 일보 그리드 — 매체별 일별 금액 + 월합계/Limit/사용율/사업계획.

    scope="month"(기본): 그 달에 거래한 매체만 — 일별 기입·예산(Limit/계획/비고)·
    자동 수집 금액이 있는 매체 + 아직 아무 이력 없는 신규 매체(기입을 시작할 수 있게).
    scope="all": 활성 매체 전부 (과거 거래처에 새로 기입할 때).
    """
    d_from, d_to = _month_range(month)
    group_meta = await _group_meta(db)
    media = sorted((await db.execute(
        select(AdMedia).where(AdMedia.active.is_(True))
    )).scalars().all(), key=_media_sort_key(group_meta))
    all_groups = _groups_payload(group_meta, media)
    budgets = {b.media_id: b for b in (await db.execute(
        select(AdMediaBudget).where(AdMediaBudget.month == month)
    )).scalars().all()}
    manual_rows = (await db.execute(
        select(AdMediaSpendDaily).where(
            AdMediaSpendDaily.date >= d_from, AdMediaSpendDaily.date < d_to
        )
    )).scalars().all()
    manual_by_media: dict = {}
    for r in manual_rows:
        manual_by_media.setdefault(r.media_id, {})[r.date.isoformat()] = r.amount

    auto_cache = {s: await _auto_daily(db, s, d_from, d_to) for s in AUTO_SOURCES}

    total_media = len(media)
    if scope == "month":
        ever_entry_ids = {mid for (mid,) in (await db.execute(
            select(AdMediaSpendDaily.media_id).distinct()
        )).all()}
        ever_budget_ids = {mid for (mid,) in (await db.execute(
            select(AdMediaBudget.media_id).distinct()
        )).all()}

        def _relevant(m: AdMedia) -> bool:
            if m.auto_source:
                return bool(auto_cache.get(m.auto_source))
            if m.id in manual_by_media or m.id in budgets:
                return True
            # 아무 월에도 이력이 없는 신규 매체 — 숨기면 추가 직후 기입을 못 하므로 노출
            return m.id not in ever_entry_ids and m.id not in ever_budget_ids

        media = [m for m in media if _relevant(m)]

    out_rows = []
    for m in media:
        daily = auto_cache.get(m.auto_source, {}) if m.auto_source else manual_by_media.get(m.id, {})
        b = budgets.get(m.id)
        total = round(sum(daily.values()), 0)
        limit_amt = b.limit_amount if b else None
        out_rows.append({
            "media_id": m.id, "name": m.name, "group_name": m.group_name,
            "inflow": m.inflow, "auto": bool(m.auto_source), "auto_source": m.auto_source,
            "owner": m.owner, "memo": m.memo,
            "daily": daily,
            "month_total": total,
            "limit_amount": limit_amt,
            "plan_amount": b.plan_amount if b else None,
            "note": b.note if b else None,  # 월별 비고
            "usage_pct": round(total / limit_amt * 100, 1) if limit_amt else None,
        })

    day_totals: dict = {}
    for r in out_rows:
        for d, v in r["daily"].items():
            day_totals[d] = day_totals.get(d, 0) + v

    grand_total = round(sum(r["month_total"] for r in out_rows), 0)
    grand_limit = round(sum(r["limit_amount"] or 0 for r in out_rows), 0)
    grand_plan = round(sum(r["plan_amount"] or 0 for r in out_rows), 0)
    return {
        "as_of": datetime.utcnow().isoformat(),
        "month": month,
        "scope": scope,
        "total_active_media": total_media,
        "days_in_month": (d_to - timedelta(days=1)).day,
        "rows": out_rows,
        "groups": all_groups,  # 판매채널 그룹 — 시트 순서·색 (rows는 이미 이 순서로 정렬됨)
        "day_totals": day_totals,
        "totals": {
            "spend": grand_total, "limit": grand_limit, "plan": grand_plan,
            "usage_pct": round(grand_total / grand_limit * 100, 1) if grand_limit else None,
        },
    }


class EntryUpsert(BaseModel):
    media_id: int
    date: str  # YYYY-MM-DD
    amount: float = Field(..., ge=0)


@router.put("/entry")
async def upsert_entry(
    payload: EntryUpsert,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    m = (await db.execute(select(AdMedia).where(AdMedia.id == payload.media_id))).scalar_one_or_none()
    if not m:
        raise HTTPException(status_code=404, detail="매체를 찾을 수 없습니다")
    if m.auto_source:
        raise HTTPException(status_code=422, detail="자동 연동 매체는 수동 기입이 불가합니다")
    try:
        d = date.fromisoformat(payload.date)
    except ValueError:
        raise HTTPException(status_code=422, detail="date는 YYYY-MM-DD 형식")
    row = (await db.execute(select(AdMediaSpendDaily).where(
        AdMediaSpendDaily.media_id == payload.media_id, AdMediaSpendDaily.date == d
    ))).scalar_one_or_none()
    if payload.amount == 0:
        if row:
            await db.delete(row)
    elif row:
        row.amount = payload.amount
    else:
        db.add(AdMediaSpendDaily(media_id=payload.media_id, date=d, amount=payload.amount))
    await db.commit()
    return {"media_id": payload.media_id, "date": payload.date, "amount": payload.amount}


class BudgetUpsert(BaseModel):
    media_id: int
    month: str = Field(..., min_length=7, max_length=7)
    limit_amount: Optional[float] = None
    plan_amount: Optional[float] = None
    note: Optional[str] = Field(None, max_length=300)  # 월별 비고


@router.put("/budget")
async def upsert_budget(
    payload: BudgetUpsert,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    row = (await db.execute(select(AdMediaBudget).where(
        AdMediaBudget.media_id == payload.media_id, AdMediaBudget.month == payload.month
    ))).scalar_one_or_none()
    if not row:
        row = AdMediaBudget(media_id=payload.media_id, month=payload.month)
        db.add(row)
    if payload.limit_amount is not None:
        row.limit_amount = payload.limit_amount
    if payload.plan_amount is not None:
        row.plan_amount = payload.plan_amount
    if payload.note is not None:
        row.note = payload.note.strip() or None
    await db.commit()
    return {"media_id": payload.media_id, "month": payload.month}


# ─── 월 단위 일괄 임포트 (시트 백필/포맷 동기화) ─────────────────────────────

def _norm_name(s: str) -> str:
    import re as _re
    return _re.sub(r"\s+", " ", (s or "")).strip()


class ImportRow(BaseModel):
    name: str = Field(..., max_length=200)
    owner: Optional[str] = Field(None, max_length=50)
    inflow: Optional[str] = Field(None, max_length=100)
    group_name: Optional[str] = Field(None, max_length=100)
    limit_amount: Optional[float] = None
    plan_amount: Optional[float] = None
    daily: dict = Field(default_factory=dict)  # {"1".."31": 금액(원)}


class MonthImport(BaseModel):
    month: str = Field(..., min_length=7, max_length=7)
    rows: list[ImportRow]
    update_master: bool = False  # True면 매체의 inflow/group/owner/sort를 이 리스트 기준으로 갱신


@router.post("/import-month")
async def import_month(
    payload: MonthImport,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """한 달치 일보를 일괄 적재 — 매체는 이름(공백 정규화)으로 매칭, 없으면 생성.

    daily 금액은 원 단위(VAT 포함). 같은 (매체,일) 기존 값은 교체.
    update_master=True면 매체 정렬/분류를 이 리스트 순서 기준으로 표준화한다.
    """
    if len(payload.rows) > 300:
        raise HTTPException(status_code=422, detail="한 번에 300행까지")
    d_from, d_to = _month_range(payload.month)
    media_all = (await db.execute(select(AdMedia))).scalars().all()
    by_name = {_norm_name(m.name): m for m in media_all}

    created = updated = entries = budgets = 0
    pending_daily: dict = {}  # (media_id, date) — 같은 요청 내 중복 방어 (동일 매체명 복수 행)
    for i, row in enumerate(payload.rows):
        key = _norm_name(row.name)
        if not key:
            continue
        m = by_name.get(key)
        if not m:
            m = AdMedia(
                name=key, owner=row.owner, inflow=row.inflow,
                group_name=row.group_name, sort_order=(i + 1) * 10,
            )
            db.add(m)
            await db.flush()
            by_name[key] = m
            created += 1
        elif payload.update_master:
            m.owner = row.owner or m.owner
            m.inflow = row.inflow or m.inflow
            m.group_name = row.group_name or m.group_name
            m.sort_order = (i + 1) * 10
            m.active = True
            updated += 1

        if row.limit_amount is not None or row.plan_amount is not None:
            b = (await db.execute(select(AdMediaBudget).where(
                AdMediaBudget.media_id == m.id, AdMediaBudget.month == payload.month
            ))).scalar_one_or_none()
            if not b:
                b = AdMediaBudget(media_id=m.id, month=payload.month)
                db.add(b)
            if row.limit_amount is not None:
                b.limit_amount = row.limit_amount
            if row.plan_amount is not None:
                b.plan_amount = row.plan_amount
            budgets += 1

        for day_str, amt in (row.daily or {}).items():
            try:
                d = date(int(payload.month[:4]), int(payload.month[5:7]), int(day_str))
            except ValueError:
                continue
            if not (d_from <= d < d_to) or amt is None:
                continue
            cache_key = (m.id, d)
            ex = pending_daily.get(cache_key) or (await db.execute(
                select(AdMediaSpendDaily).where(
                    AdMediaSpendDaily.media_id == m.id, AdMediaSpendDaily.date == d
                )
            )).scalar_one_or_none()
            if float(amt) == 0:
                if ex:
                    await db.delete(ex)
                    pending_daily.pop(cache_key, None)
                continue
            if ex:
                ex.amount = float(amt)
            else:
                ex = AdMediaSpendDaily(media_id=m.id, date=d, amount=float(amt))
                db.add(ex)
            pending_daily[cache_key] = ex
            entries += 1

    await db.commit()
    return {"month": payload.month, "media_created": created, "media_updated": updated,
            "entries": entries, "budgets": budgets}


# ─── 엑셀 업로드 (시트 xlsx 그대로 → 미리보기 → 반영) ─────────────────────────
# 팀 구글시트를 xlsx로 받아 올리면 그 달 일보가 시트와 같은 숫자가 되게 한다.
# 파싱·자동 연결은 services/adspend_xlsx(순수 함수). 미리보기는 DB를 바꾸지 않고,
# 반영은 같은 파일을 다시 읽어 사람이 고른 연결대로 한 트랜잭션에 쓴다(파일은 남기지 않음).

XLSX_MAX_BYTES = 8 * 1024 * 1024


def _check_month(month: str) -> str:
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month or ""):
        raise HTTPException(status_code=422, detail="month는 YYYY-MM 형식")
    return month


async def _parse_upload(file: UploadFile, month: str, sheet: Optional[str]) -> xl.ParsedSheet:
    data = await file.read(XLSX_MAX_BYTES + 1)
    if len(data) > XLSX_MAX_BYTES:
        raise HTTPException(status_code=413, detail="파일이 너무 큽니다 (8MB까지)")
    if data[:2] != b"PK":
        raise HTTPException(status_code=422, detail="엑셀(.xlsx) 파일만 올릴 수 있습니다 — 구글시트에서 파일 > 다운로드 > Microsoft Excel(.xlsx)로 받아 주세요")
    try:  # openpyxl 파싱은 CPU 작업 — 이벤트 루프(수집 스케줄러 등)를 막지 않게 스레드로
        return await asyncio.to_thread(xl.parse_workbook, data, month, (sheet or "").strip() or None)
    except xl.XlsxImportError as e:  # 시트를 못 고르면 탭 목록도 줘서 화면이 고르게 한다
        raise HTTPException(status_code=422, detail={"message": str(e), "sheets": e.sheets} if e.sheets else str(e))


async def _load_aliases(db: AsyncSession) -> list:
    """기억된 연결 — 테이블 문제로 업로드 전체가 막히지 않게 실패하면 빈 목록."""
    try:
        return list((await db.execute(select(AdMediaAlias))).scalars().all())
    except Exception as e:
        logger.warning("ad_media_aliases 읽기 실패 — 기억된 연결 없이 진행: %s", e)
        await db.rollback()
        return []


async def _xlsx_media(db: AsyncSession, month: str, current_user: User):
    """자동 연결 재료 — (MediaInfo 목록(비활성 포함), 이 달 보드(활성 매체 전부), {id: AdMedia}, 별칭 목록)."""
    aliases = await _load_aliases(db)
    board = await spend_board(month=month, scope="all", current_user=current_user, db=db)
    brow = {r["media_id"]: r for r in board["rows"]}
    by_media: dict = {}
    for a in aliases:
        by_media.setdefault(a.media_id, set()).add(a.key)
    rows = (await db.execute(select(AdMedia))).scalars().all()
    infos = [xl.MediaInfo(
        id=m.id, name=m.name, inflow=m.inflow, group=m.group_name, active=bool(m.active),
        auto=bool(m.auto_source), month_spend=float((brow.get(m.id) or {}).get("month_total") or 0),
        month_limit=(brow.get(m.id) or {}).get("limit_amount"), aliases=by_media.get(m.id, set()),
    ) for m in rows]
    return infos, board, {m.id: m for m in rows}, aliases


@router.post("/import-xlsx/preview")
async def import_xlsx_preview(
    file: UploadFile = File(...),
    month: str = Form(...),
    sheet: Optional[str] = Form(None),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """시트 xlsx 미리보기 — DB는 바꾸지 않는다.

    행별 자동 연결(+후보)과, 화면이 반영 후 합계를 계산할 재료(매체별 이 달 금액·현재 보드 합계)를 준다.
    """
    month = _check_month(month)
    parsed = await _parse_upload(file, month, sheet)
    media, board, _, _ = await _xlsx_media(db, month, current_user)
    sug = xl.suggest(parsed.rows, media)
    order = {r["media_id"]: i for i, r in enumerate(board["rows"])}  # 보드(시트) 순서, 숨긴 매체는 뒤
    media_sorted = sorted(media, key=lambda m: (order.get(m.id, 10 ** 6), m.inflow or "", m.name))
    return {
        "month": month,
        "file": {
            "name": file.filename, "sheet": parsed.sheet, "sheets": parsed.sheets,
            "detected_month": parsed.detected_month, "header_row": parsed.header_row,
            "total_row": parsed.total_row, "has_note_col": parsed.has_note_col,
        },
        "official": {"spend": parsed.official_total, "limit": parsed.official_limit},
        "computed": {"spend": round(sum(r.total for r in parsed.rows), 2),
                     "limit": round(sum(r.limit or 0 for r in parsed.rows), 2)},
        "warnings": parsed.warnings,
        "excluded": parsed.excluded,
        "rows": [{
            "key": r.key, "row": r.row, "inflow": r.inflow, "group": r.group, "name": r.name,
            "owner": r.owner, "total": round(r.total, 2), "limit": r.limit, "note": r.note,
            "days": len(r.amounts), "flags": r.flags, "nonnum": r.nonnum,
            "suggestion": sug[r.key].as_dict(), "candidates": xl.candidates(r, media),
        } for r in parsed.rows],
        "media": [{
            "id": m.id, "name": m.name, "inflow": m.inflow, "group_name": m.group, "active": m.active,
            "auto": m.auto, "month_spend": m.month_spend, "month_limit": m.month_limit,
        } for m in media_sorted],
        "board": board["totals"],
    }


@router.post("/import-xlsx/apply")
async def import_xlsx_apply(
    file: UploadFile = File(...),
    month: str = Form(...),
    sheet: Optional[str] = Form(None),
    mapping: str = Form("{}"),  # {"r57": 114 | "new" | "skip"} — 빠진 행은 자동 연결대로
    clear_missing: bool = Form(True),  # 파일에 없는 매체의 이 달 금액·Limit 비움
    remember: bool = Form(True),  # 이름이 다른 연결을 기억해 다음 달부터 자동으로
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """시트 xlsx 반영 — 연결된 매체의 그 달 일별 금액·Limit를 파일 값으로 교체한다.

    새 매체는 판매채널 그룹 맨 아래에 만들고, 숨긴 매체로 연결하면 다시 켠다.
    자동 수집 매체(메타·네이버 검색광고)에는 넣지 않는다. 사업계획 금액은 건드리지 않는다.
    """
    month = _check_month(month)
    parsed = await _parse_upload(file, month, sheet)
    media, _, by_id, aliases = await _xlsx_media(db, month, current_user)
    sug = xl.suggest(parsed.rows, media)
    try:
        chosen = json.loads(mapping or "{}")
    except ValueError:
        chosen = None
    if not isinstance(chosen, dict):
        raise HTTPException(status_code=422, detail="mapping은 {행키: 매체id|'new'|'skip'} JSON")
    unknown = [k for k in chosen if k not in {r.key for r in parsed.rows}]
    if unknown:
        raise HTTPException(status_code=422, detail=f"파일에 없는 행 {unknown[:5]} — 미리보기 뒤 파일이 바뀌었으면 다시 올려 주세요")

    plan: dict = {}  # row.key → 매체 id | "new" | "skip"
    for r in parsed.rows:
        v = chosen.get(r.key)
        if v is None:
            s = sug[r.key]
            v = s.media_id if s.action == "map" else s.action
        if v in ("new", "skip"):
            plan[r.key] = v
            continue
        try:
            m = by_id.get(int(v))
        except (TypeError, ValueError):
            m = None
        if m is None:
            raise HTTPException(status_code=422, detail=f"'{r.name}' 행: 없는 매체입니다 ({v})")
        if m.auto_source:
            raise HTTPException(status_code=422, detail=f"'{m.name}'은 자동 수집 매체라 엑셀 금액을 넣을 수 없습니다 — 다른 매체를 고르거나 건너뛰기로 두세요")
        plan[r.key] = m.id

    d_from, d_to = _month_range(month)
    y, mo = int(month[:4]), int(month[5:7])

    # 새 매체 — 같은 (판매채널, 매체명) 행은 한 매체로. 그룹 맨 아래, 새 그룹은 그룹 목록 맨 끝
    group_meta = await _group_meta(db)
    gmax = max([g.sort_order or 0 for g in group_meta.values()] or [0])
    next_sort: dict = {}
    new_media: dict = {}  # alias_key → AdMedia
    for r in parsed.rows:
        k = xl.alias_key(r.inflow, r.name)
        if plan[r.key] != "new" or k in new_media:
            continue
        inflow = _norm_name(r.inflow)[:100] or None
        if inflow not in next_sort:
            q = select(func.max(AdMedia.sort_order)).where(AdMedia.active.is_(True))
            q = q.where(AdMedia.inflow == inflow) if inflow else q.where(AdMedia.inflow.is_(None))
            next_sort[inflow] = int((await db.execute(q)).scalar() or 0) + 10
        if inflow and inflow not in group_meta:
            gmax += 10
            group_meta[inflow] = AdMediaGroup(name=inflow, sort_order=gmax)
            db.add(group_meta[inflow])
        new_media[k] = AdMedia(
            name=_norm_name(r.name)[:200], inflow=inflow, group_name=_norm_name(r.group)[:100] or None,
            owner=_norm_name(r.owner)[:50] or None, sort_order=next_sort[inflow], active=True,
        )
        next_sort[inflow] += 10
        db.add(new_media[k])
    if new_media:
        await db.flush()

    targets: dict = {}  # 매체 id → [행]
    for r in parsed.rows:
        v = plan[r.key]
        if v != "skip":
            mid = new_media[xl.alias_key(r.inflow, r.name)].id if v == "new" else v
            targets.setdefault(mid, []).append(r)

    reactivated = 0
    for mid in targets:
        m = by_id.get(mid)
        if m is not None and not m.active:
            m.active = True
            reactivated += 1

    cleared = [m.id for m in media if m.active and not m.auto and m.has_month
               and m.id not in targets] if clear_missing else []
    touched = list(targets) + cleared
    if touched:
        await db.execute(delete(AdMediaSpendDaily).where(
            AdMediaSpendDaily.media_id.in_(touched),
            AdMediaSpendDaily.date >= d_from, AdMediaSpendDaily.date < d_to,
        ))

    entries = 0
    for mid, rs in targets.items():
        daily: dict = {}
        for r in rs:
            for d, v in r.amounts.items():
                daily[d] = daily.get(d, 0.0) + v
        for d, v in sorted(daily.items()):
            if round(v, 2):
                db.add(AdMediaSpendDaily(media_id=mid, date=date(y, mo, d), amount=round(v, 2)))
                entries += 1

    # Limit = 연결된 행 Limit 합 (없으면 비움). 비고는 파일에 비고 칸이 있을 때만 교체
    budgets = {b.media_id: b for b in (await db.execute(
        select(AdMediaBudget).where(AdMediaBudget.month == month, AdMediaBudget.media_id.in_(touched or [-1]))
    )).scalars().all()}
    for mid in touched:
        rs = targets.get(mid, [])
        lims = [r.limit for r in rs if r.limit is not None]
        limit = round(sum(lims), 2) if lims else None
        note = "; ".join(dict.fromkeys(r.note for r in rs if r.note))[:300] or None
        b = budgets.get(mid)
        if b is None:
            if limit is None and not (parsed.has_note_col and note):
                continue
            b = AdMediaBudget(media_id=mid, month=month)
            db.add(b)
        b.limit_amount = limit
        if parsed.has_note_col and mid in targets:
            b.note = note
        if b.id is not None and b.limit_amount is None and b.plan_amount is None and not b.note:
            await db.delete(b)
    await db.commit()

    # 연결 기억 — 금액 반영과 따로 커밋 (실패해도 반영은 유지)
    remembered = 0
    remember_error = None
    if remember:
        try:
            existing = {a.key: a for a in aliases}
            for r in parsed.rows:
                v = plan[r.key]
                if v == "skip":
                    continue
                k = xl.alias_key(r.inflow, r.name)
                m = new_media[k] if v == "new" else by_id[v]
                a = existing.get(k)
                if a is None:
                    if xl.plain_match(r, m.name, m.inflow):
                        continue  # 이름·판매채널이 같으면 기억할 필요 없음
                    a = AdMediaAlias(key=k, media_id=m.id, inflow=r.inflow[:100] or None, name=r.name[:200])
                    db.add(a)
                    existing[k] = a
                elif a.media_id != m.id:
                    a.media_id = m.id
                else:
                    continue
                remembered += 1
            await db.commit()
        except Exception as e:
            logger.warning("ad_media_aliases 저장 실패: %s", e)
            await db.rollback()
            remembered, remember_error = 0, "연결 기억 저장에 실패했습니다 (금액 반영은 완료)"

    db.expire_all()
    board = await spend_board(month=month, scope="all", current_user=current_user, db=db)
    return {
        "month": month, "sheet": parsed.sheet,
        "rows": len(parsed.rows),
        "mapped": sum(1 for v in plan.values() if v != "skip"),
        "skipped": sum(1 for v in plan.values() if v == "skip"),
        "media_created": len(new_media), "media_reactivated": reactivated, "media_cleared": len(cleared),
        "entries": entries, "remembered": remembered, "remember_error": remember_error,
        "official": {"spend": parsed.official_total, "limit": parsed.official_limit},
        "board": board["totals"],
    }


@router.get("/export")
async def export_month_xlsx(
    month: str = Query(..., min_length=7, max_length=7),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """일보 엑셀 다운로드 — 10월 포맷(담당자|판매채널|광고분류|광고매체|일별|TOTAL|계획|사용율)."""
    from io import BytesIO
    from urllib.parse import quote

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    board = await spend_board(month=month, current_user=current_user, db=db)
    media_meta = {m.id: m for m in (await db.execute(select(AdMedia))).scalars().all()}
    days = board["days_in_month"]

    wb = Workbook()
    ws = wb.active
    ws.title = f"광고비 일보 {int(month[5:7])}월"
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    money = "#,##0"

    headers = ["판매 담당자", "판매채널", "광고분류", "광고매체"] + \
        [f"{int(month[5:7])}.{d}" for d in range(1, days + 1)] + \
        ["TOTAL", "광고비 계획(Limit)", "사용율", "비고"]
    for ci, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    group_color = {g["name"]: g["color"] for g in board.get("groups", []) if g.get("color")}
    ri = 2
    for r in board["rows"]:
        meta = media_meta.get(r["media_id"])
        ws.cell(ri, 1, (meta.owner if meta else None) or "")
        ch = ws.cell(ri, 2, r["inflow"] or "")
        gc = group_color.get(r["inflow"])
        if gc:
            ch.fill = PatternFill(start_color=gc[1:], end_color=gc[1:], fill_type="solid")
        ws.cell(ri, 3, r["group_name"] or "")
        ws.cell(ri, 4, r["name"] + (" (자동)" if r["auto"] else ""))
        for d in range(1, days + 1):
            v = r["daily"].get(f"{month}-{d:02d}")
            if v:
                c = ws.cell(ri, 4 + d, round(v))
                c.number_format = money
        ws.cell(ri, 5 + days, round(r["month_total"])).number_format = money
        if r["limit_amount"]:
            ws.cell(ri, 6 + days, round(r["limit_amount"])).number_format = money
        if r["usage_pct"] is not None:
            ws.cell(ri, 7 + days, r["usage_pct"] / 100).number_format = "0%"
        ws.cell(ri, 8 + days, r.get("note") or (meta.memo if meta else None) or "")
        ri += 1

    # 합계 행
    ws.cell(ri, 4, "합계").font = header_font
    for d in range(1, days + 1):
        v = board["day_totals"].get(f"{month}-{d:02d}")
        if v:
            c = ws.cell(ri, 4 + d, round(v))
            c.number_format = money
            c.font = header_font
    ws.cell(ri, 5 + days, round(board["totals"]["spend"])).number_format = money
    ws.cell(ri, 5 + days).font = header_font
    if board["totals"]["limit"]:
        ws.cell(ri, 6 + days, round(board["totals"]["limit"])).number_format = money
    if board["totals"]["usage_pct"] is not None:
        ws.cell(ri, 7 + days, board["totals"]["usage_pct"] / 100).number_format = "0%"

    ws.freeze_panes = "E2"
    for i, w in enumerate([10, 14, 12, 30] + [9] * days + [13, 14, 8, 16], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.cell(ri + 2, 4, "단위: 원 (VAT 포함) · 시스템 광고비 일보 export")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"광고비일보_{month}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# ─── SALES(CSA) 채널 매출 연동 ───────────────────────────────────────────────
# myfirstproject CSA 대시보드에서 월×채널 매출(공급가, VAT별도·월정액 차감 후)을
# 가져와 ROAS를 자동 계산한다. 이름이 명확히 대응하는 채널만 매핑하고(대표 지시),
# 애매한 채널(B2B·오프라인 CVS/마트·GS샵 홈쇼핑·11번가 슈팅 등)은 수동 기입 유지.
import time as _time

CSA_BASE = "https://myfirstproject-api-557811875995.asia-northeast3.run.app"

INFLOW_TO_CSA: dict = {
    "자사몰": ["카페24"],
    "네이버스토어": ["스마트스토어"],
    "쿠팡 WING": ["쿠팡 WING"],
    "쿠팡 (사입)": ["쿠팡 로켓"],
    "토스쇼핑": ["토스"],
    "G마켓": ["지마켓"],
    "옥션": ["옥션"],
    "알리익스프레스": ["알리익스프레스"],
    "11번가": ["11번가"],
    "카카오 스토어": ["카카오톡스토어"],
    "카카오선물하기": ["카카오선물하기"],
    "올리브영": ["올리브영"],
    "올웨이즈": ["올웨이즈"],
    "GS샵": ["GS 샵"],
    "CJ온스타일": ["CJ온스타일"],
    "롯데ON": ["롯데온"],
    "지그재그": ["카카오스타일"],  # 지그재그 운영명 = 카카오스타일
    "이지웰": ["이지웰"],
    "에이블리": ["에이블리"],
    "Tdeal": ["T deal"],
    "신세계라이브쇼핑": ["신세계 라이브쇼핑"],
    "신세계TV": ["신세계 TV 쇼핑"],
    "비마트": ["B마트"],
    "컬리": ["마켓컬리"],
}

_csa_cache: dict = {}  # month -> {"data": {channel_name: revenue}, "expires": ts}
_CSA_TTL = 600


async def _csa_month_revenues(month: str) -> Optional[dict]:
    """해당 월의 CSA 채널별 매출 {채널명: revenue}. 실패 시 None (수동 기입 폴백)."""
    now = _time.time()
    cached = _csa_cache.get(month)
    if cached and cached["expires"] > now:
        return cached["data"]
    d_from, d_to = _month_range(month)
    try:
        # SCM API는 2026-10-24부터 인증 필수 — 서비스 키(SCM JARVIS_API_KEY와 같은 값)를 env SCM_API_KEY로 주입
        _scm_key = os.getenv("SCM_API_KEY", "")
        _headers = {"Authorization": f"Bearer {_scm_key}"} if _scm_key else {}
        async with httpx.AsyncClient(timeout=30) as client:
            resp = await client.get(f"{CSA_BASE}/api/csa/dashboard", headers=_headers, params={
                "period_start": d_from.isoformat(),
                "period_end": (d_to - timedelta(days=1)).isoformat(),
                "granularity": "month",
            })
        if resp.status_code != 200:
            logger.warning(f"[CSA] dashboard {month} -> {resp.status_code}")
            return None
        data = {c["channel_name"]: float(c.get("revenue") or 0)
                for c in resp.json().get("channels", [])}
        _csa_cache[month] = {"data": data, "expires": now + _CSA_TTL}
        return data
    except Exception as e:
        logger.warning(f"[CSA] dashboard {month} 실패: {e}")
        return None


# ─── 채널 ROAS 보드 (유입채널 × 월 — 광고비 자동 + 매출 기입) ────────────────

async def _inflow_spend_by_month(db: AsyncSession, months: list) -> dict:
    """{(month, inflow): spend} — 일보 수동 기입 + 자동 매체(meta/naver_sa)."""
    m_from, m_to = months[0], months[-1]
    d_from, _ = _month_range(m_from)
    _, d_to = _month_range(m_to)

    media = (await db.execute(select(AdMedia).where(AdMedia.active.is_(True)))).scalars().all()
    inflow_of = {m.id: (m.inflow or "기타") for m in media}
    auto_media = [(m.id, m.auto_source) for m in media if m.auto_source]

    out: dict = {}
    rows = (await db.execute(
        select(
            AdMediaSpendDaily.media_id,
            func.to_char(AdMediaSpendDaily.date, "YYYY-MM"),
            func.coalesce(func.sum(AdMediaSpendDaily.amount), 0),
        ).where(AdMediaSpendDaily.date >= d_from, AdMediaSpendDaily.date < d_to)
        .group_by(text("1, 2"))  # to_char 포맷이 바인드 파라미터로 렌더돼 표현식 GROUP BY가 불일치 → 위치 지정
    )).all()
    for mid, month, amt in rows:
        if mid not in inflow_of:
            continue  # 비활성(삭제) 매체 지출은 보드와 동일하게 제외 (홈 카드 ↔ 보드 정합)
        key = (month, inflow_of[mid])
        out[key] = out.get(key, 0) + float(amt)
    for mid, source in auto_media:
        daily = await _auto_daily(db, source, d_from, d_to)
        for d, amt in daily.items():
            key = (d[:7], inflow_of.get(mid, "기타"))
            out[key] = out.get(key, 0) + amt
    return out


@router.get("/roas-board")
async def roas_board(
    months_back: int = Query(default=6, ge=1, le=24),
    month_from: str | None = Query(default=None, min_length=7, max_length=7),
    month_to: str | None = Query(default=None, min_length=7, max_length=7),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """유입채널별 월 광고비(일보 자동 합산) × 매출(기입/자사몰 자동) × ROAS.

    구 '채널 성과 분석'(monthly_channel_spends 소수 채널) 대체 —
    광고비 일보의 전체 유입채널 축으로 분석한다. 금액 원 단위(VAT 포함).
    month_from/month_to(YYYY-MM)를 함께 주면 커스텀 범위(최대 24개월)가 months_back보다 우선.
    """
    from app.models import ChannelRevenue, MallOrder

    this_month = date.today().strftime("%Y-%m")
    if month_from and month_to:
        import re as _re
        if not (_re.fullmatch(r"\d{4}-\d{2}", month_from) and _re.fullmatch(r"\d{4}-\d{2}", month_to)):
            raise HTTPException(status_code=422, detail="month_from/month_to는 YYYY-MM 형식이어야 합니다")
        if month_from > month_to:
            month_from, month_to = month_to, month_from
        span = (int(month_to[:4]) - int(month_from[:4])) * 12 + int(month_to[5:7]) - int(month_from[5:7]) + 1
        if span > 24:
            raise HTTPException(status_code=422, detail="조회 범위는 최대 24개월입니다")
        months = [_month_add(month_from, i) for i in range(span)]
    else:
        months = [_month_add(this_month, i) for i in range(-months_back + 1, 1)]
    spend_by = await _inflow_spend_by_month(db, months)

    # 매출: ① SALES(CSA) 매칭 채널 = 자동 ② 미매칭 채널 = 수동 기입(channel_revenues)
    import asyncio as _asyncio
    rev_rows = (await db.execute(
        select(ChannelRevenue).where(
            ChannelRevenue.month >= months[0], ChannelRevenue.month <= months[-1]
        )
    )).scalars().all()
    rev_by = {(r.month, r.inflow): r.revenue for r in rev_rows}
    csa_results = await _asyncio.gather(*[_csa_month_revenues(m) for m in months])
    csa_by_month = dict(zip(months, csa_results))

    # 예산(Limit·사업계획) — 유입채널별 월합
    budgets = (await db.execute(
        select(AdMediaBudget).where(AdMediaBudget.month.in_(months))
    )).scalars().all()
    media = (await db.execute(select(AdMedia))).scalars().all()
    inflow_of = {m.id: (m.inflow or "기타") for m in media}
    limit_by: dict = {}
    for b in budgets:
        key = (b.month, inflow_of.get(b.media_id, "기타"))
        limit_by[key] = limit_by.get(key, 0) + (b.limit_amount or 0)

    inflows = sorted({k[1] for k in list(spend_by) + list(rev_by) + list(limit_by)})
    cells = []
    for m in months:
        csa = csa_by_month.get(m)
        for inf in inflows:
            spend = round(spend_by.get((m, inf), 0), 0)
            csa_names = INFLOW_TO_CSA.get(inf)
            revenue = None
            source = None
            if csa_names and csa is not None:
                revenue = sum(csa.get(n, 0) for n in csa_names)
                source = "sales"
                if revenue == 0:
                    revenue = None  # 해당 월 SALES 데이터 없음(미업로드) — 빈 값으로
            if revenue is None and rev_by.get((m, inf)) is not None:
                revenue = rev_by[(m, inf)]
                source = "manual"
            limit_amt = round(limit_by.get((m, inf), 0), 0) or None
            if not spend and not revenue and not limit_amt:
                continue
            cells.append({
                "month": m, "inflow": inf,
                "spend": spend,
                "limit": limit_amt,
                "usage_pct": round(spend / limit_amt * 100, 1) if limit_amt else None,
                "revenue": round(revenue, 0) if revenue is not None else None,
                "revenue_auto": source == "sales",
                "revenue_source": source,
                "sales_linked": bool(csa_names),
                "roas": round(revenue / spend, 2) if revenue and spend else None,
            })
    return {
        "as_of": datetime.utcnow().isoformat(),
        "months": months,
        "this_month": this_month,
        "inflows": inflows,
        "cells": cells,
    }


class RevenueUpsert(BaseModel):
    month: str = Field(..., min_length=7, max_length=7)
    inflow: str = Field(..., max_length=100)
    revenue: float = Field(..., ge=0)


@router.put("/revenue")
async def upsert_channel_revenue(
    payload: RevenueUpsert,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    from app.models import ChannelRevenue

    if payload.inflow in INFLOW_TO_CSA:
        raise HTTPException(status_code=422, detail="SALES 시스템 연동 채널은 매출이 자동 집계됩니다")
    row = (await db.execute(select(ChannelRevenue).where(
        ChannelRevenue.month == payload.month, ChannelRevenue.inflow == payload.inflow
    ))).scalar_one_or_none()
    if payload.revenue == 0:
        if row:
            await db.delete(row)
    elif row:
        row.revenue = payload.revenue
    else:
        db.add(ChannelRevenue(month=payload.month, inflow=payload.inflow, revenue=payload.revenue))
    await db.commit()
    return {"month": payload.month, "inflow": payload.inflow, "revenue": payload.revenue}


@router.get("/trend")
async def adspend_trend(
    month_from: str | None = Query(default=None, min_length=7, max_length=7),
    month_to: str | None = Query(default=None, min_length=7, max_length=7),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """비용 분석 그래프용 — 기간 내 월별 집행(유입채널 스택)/Limit/사업계획 시계열.

    기본 기간: 이번 달 포함 최근 6개월. 최대 24개월.
    """
    today = date.today()
    end = month_to or f"{today.year:04d}-{today.month:02d}"
    start = month_from or _month_add(end, -5)
    if start > end:
        start, end = end, start
    months = [start]
    while months[-1] < end and len(months) < 24:
        months.append(_month_add(months[-1], 1))

    spend_by = await _inflow_spend_by_month(db, months)
    budget_rows = (await db.execute(
        select(
            AdMediaBudget.month,
            func.coalesce(func.sum(AdMediaBudget.limit_amount), 0),
            func.coalesce(func.sum(AdMediaBudget.plan_amount), 0),
        ).where(AdMediaBudget.month.in_(months)).group_by(AdMediaBudget.month)
    )).all()
    budget_of = {m: {"limit": float(l), "plan": float(p)} for m, l, p in budget_rows}

    inflows = sorted({inf for (_m, inf) in spend_by.keys()})
    series = []
    for m in months:
        by_inflow = {inf: round(spend_by.get((m, inf), 0), 0) for inf in inflows}
        b = budget_of.get(m, {"limit": 0, "plan": 0})
        total = round(sum(by_inflow.values()), 0)
        series.append({
            "month": m,
            "total": total,
            "limit": round(b["limit"], 0),
            "plan": round(b["plan"], 0),
            "usage_pct": round(total / b["limit"] * 100, 1) if b["limit"] else None,
            "by_inflow": by_inflow,
        })
    return {"months": months, "inflows": inflows, "series": series}


@router.get("/monthly-summary")
async def monthly_summary(
    month: str = Query(..., min_length=7, max_length=7),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """KPI 요약용 — 유입채널(inflow) 그룹별 집행/Limit/사업계획 합계."""
    board = await spend_board(month=month, current_user=current_user, db=db)
    groups: dict = {}
    for r in board["rows"]:
        g = groups.setdefault(r["inflow"] or "기타", {"spend": 0, "limit": 0, "plan": 0})
        g["spend"] += r["month_total"]
        g["limit"] += r["limit_amount"] or 0
        g["plan"] += r["plan_amount"] or 0
    return {
        "month": month,
        "by_inflow": [{"inflow": k, **{kk: round(vv, 0) for kk, vv in v.items()}} for k, v in groups.items()],
        "totals": board["totals"],
    }
