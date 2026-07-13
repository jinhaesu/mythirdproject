"""인플루언서 시딩 관리 모듈 — 시딩 비용/대상 기록 + AI 타겟 고객층 분석 + 집계.

엔드포인트:
  GET    /api/v1/influencer/seedings                — 시딩 목록
  POST   /api/v1/influencer/seedings                — 시딩 등록
  PUT    /api/v1/influencer/seedings/{id}            — 시딩 수정
  DELETE /api/v1/influencer/seedings/{id}            — 시딩 삭제
  POST   /api/v1/influencer/seedings/{id}/analyze    — AI 타겟 고객층 분석
  GET    /api/v1/influencer/summary                  — 채널/타겟/월별 집계
  GET    /api/v1/influencer/export                   — xlsx 다운로드
"""
import logging
import re
from collections import defaultdict
from datetime import date, datetime
from typing import Optional

import httpx
from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.v1.endpoints.auth import get_current_user
from app.db.database import get_db
from app.models.influencer import InfluencerSeeding
from app.models.user import User

logger = logging.getLogger(__name__)
router = APIRouter()

_UNANALYZED_LABEL = "미분석"


# ── 공용 헬퍼 ────────────────────────────────────────────────────────────────

def _serialize(row: InfluencerSeeding) -> dict:
    return {
        "id": row.id,
        "name": row.name,
        "channel": row.channel,
        "url": row.url,
        "follower_count": row.follower_count,
        "cost": row.cost,
        "seeded_at": row.seeded_at.isoformat() if row.seeded_at else None,
        "product": row.product,
        "notes": row.notes,
        "ai_target_segment": row.ai_target_segment,
        "ai_audience_summary": row.ai_audience_summary,
        "ai_analyzed_at": row.ai_analyzed_at.isoformat() if row.ai_analyzed_at else None,
        "created_at": row.created_at.isoformat() if row.created_at else None,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
    }


def _parse_date(value: str, field_name: str = "seeded_at") -> date:
    try:
        return date.fromisoformat(value)
    except (ValueError, TypeError):
        raise HTTPException(status_code=422, detail=f"{field_name}는 YYYY-MM-DD 형식이어야 합니다.")


async def _fetch_page_text(url: str) -> Optional[str]:
    """URL의 title/meta description/본문 앞부분을 간단히 추출. 실패해도 None 반환."""
    try:
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
            )
        }
        async with httpx.AsyncClient(timeout=15.0, follow_redirects=True) as client:
            resp = await client.get(url, headers=headers)
        if resp.status_code >= 400:
            return None
        html = resp.text
    except Exception as e:
        logger.warning(f"[Influencer] 페이지 fetch 실패 url={url}: {e}")
        return None

    parts = []
    title_m = re.search(r"<title[^>]*>(.*?)</title>", html, re.IGNORECASE | re.DOTALL)
    if title_m:
        parts.append(f"[title] {_strip_html(title_m.group(1))}")

    desc_m = re.search(
        r'<meta[^>]+name=["\']description["\'][^>]+content=["\'](.*?)["\']',
        html,
        re.IGNORECASE | re.DOTALL,
    )
    if desc_m:
        parts.append(f"[description] {_strip_html(desc_m.group(1))}")

    og_m = re.search(
        r'<meta[^>]+property=["\']og:description["\'][^>]+content=["\'](.*?)["\']',
        html,
        re.IGNORECASE | re.DOTALL,
    )
    if og_m:
        parts.append(f"[og:description] {_strip_html(og_m.group(1))}")

    # 본문 앞부분: 스크립트/스타일 제거 후 태그 스트립
    body = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.IGNORECASE)
    body = re.sub(r"<style[\s\S]*?</style>", " ", body, flags=re.IGNORECASE)
    body_text = _strip_html(body)
    if body_text:
        parts.append(f"[body] {body_text[:3000]}")

    combined = "\n".join(parts).strip()
    return combined[:3500] if combined else None


def _strip_html(text: str) -> str:
    text = re.sub(r"<[^>]+>", " ", text or "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


# ── GET /seedings ────────────────────────────────────────────────────────────

@router.get("/seedings")
async def list_seedings(
    channel: Optional[str] = Query(default=None),
    since: Optional[str] = Query(default=None, description="YYYY-MM-DD"),
    until: Optional[str] = Query(default=None, description="YYYY-MM-DD"),
    q: Optional[str] = Query(default=None, description="이름/제품 검색어"),
    limit: int = Query(default=200, ge=1, le=1000),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """인플루언서 시딩 목록 (seeded_at desc)."""
    query = select(InfluencerSeeding)
    if channel:
        query = query.where(InfluencerSeeding.channel == channel)
    if since:
        query = query.where(InfluencerSeeding.seeded_at >= _parse_date(since, "since"))
    if until:
        query = query.where(InfluencerSeeding.seeded_at <= _parse_date(until, "until"))
    if q:
        like = f"%{q.strip()}%"
        query = query.where(
            (InfluencerSeeding.name.ilike(like)) | (InfluencerSeeding.product.ilike(like))
        )
    query = query.order_by(InfluencerSeeding.seeded_at.desc(), InfluencerSeeding.id.desc()).limit(limit)

    rows = (await db.execute(query)).scalars().all()
    return {"seedings": [_serialize(r) for r in rows], "count": len(rows)}


# ── POST /seedings ───────────────────────────────────────────────────────────

class InfluencerSeedingCreate(BaseModel):
    name: str
    channel: str
    url: Optional[str] = None
    follower_count: Optional[int] = None
    cost: Optional[float] = 0
    seeded_at: str  # YYYY-MM-DD
    product: Optional[str] = None
    notes: Optional[str] = None


@router.post("/seedings")
async def create_seeding(
    payload: InfluencerSeedingCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """인플루언서 시딩 등록."""
    if not payload.name or not payload.name.strip():
        raise HTTPException(status_code=422, detail="name은 필수입니다.")
    if not payload.channel or not payload.channel.strip():
        raise HTTPException(status_code=422, detail="channel은 필수입니다.")

    seeded_at = _parse_date(payload.seeded_at)

    row = InfluencerSeeding(
        name=payload.name.strip(),
        channel=payload.channel.strip(),
        url=payload.url,
        follower_count=payload.follower_count,
        cost=payload.cost or 0,
        seeded_at=seeded_at,
        product=payload.product,
        notes=payload.notes,
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return _serialize(row)


# ── PUT /seedings/{id} ───────────────────────────────────────────────────────

class InfluencerSeedingUpdate(BaseModel):
    name: Optional[str] = None
    channel: Optional[str] = None
    url: Optional[str] = None
    follower_count: Optional[int] = None
    cost: Optional[float] = None
    seeded_at: Optional[str] = None
    product: Optional[str] = None
    notes: Optional[str] = None


@router.put("/seedings/{seeding_id}")
async def update_seeding(
    seeding_id: int,
    payload: InfluencerSeedingUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """인플루언서 시딩 부분 수정 (지정된 필드만 갱신)."""
    row = (
        await db.execute(select(InfluencerSeeding).where(InfluencerSeeding.id == seeding_id))
    ).scalar_one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="시딩 항목을 찾을 수 없습니다.")

    data = payload.model_dump(exclude_none=True)
    if "seeded_at" in data:
        data["seeded_at"] = _parse_date(data["seeded_at"])

    for field, value in data.items():
        setattr(row, field, value)

    await db.commit()
    await db.refresh(row)
    return _serialize(row)


# ── DELETE /seedings/{id} ────────────────────────────────────────────────────

@router.delete("/seedings/{seeding_id}")
async def delete_seeding(
    seeding_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    row = (
        await db.execute(select(InfluencerSeeding).where(InfluencerSeeding.id == seeding_id))
    ).scalar_one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="시딩 항목을 찾을 수 없습니다.")
    await db.delete(row)
    await db.commit()
    return {"status": "deleted", "id": seeding_id}


# ── POST /seedings/{id}/analyze ──────────────────────────────────────────────

@router.post("/seedings/{seeding_id}/analyze")
async def analyze_seeding(
    seeding_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """URL 페이지 텍스트 + 메타데이터를 바탕으로 AI가 타겟 고객층을 분석."""
    row = (
        await db.execute(select(InfluencerSeeding).where(InfluencerSeeding.id == seeding_id))
    ).scalar_one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="시딩 항목을 찾을 수 없습니다.")

    page_text: Optional[str] = None
    if row.url:
        page_text = await _fetch_page_text(row.url)

    from app.services.ai import ClaudeService

    claude = ClaudeService()
    try:
        result = await claude.analyze_influencer_target(
            name=row.name,
            channel=row.channel,
            url=row.url,
            page_text=page_text,
            follower_count=row.follower_count,
            product=row.product,
            notes=row.notes,
        )
    except Exception as e:
        logger.error(f"[Influencer] AI 분석 실패 id={seeding_id}: {e}", exc_info=True)
        raise HTTPException(status_code=502, detail=f"AI 분석 요청 실패: {e}")

    row.ai_target_segment = result.get("target_segment")
    row.ai_audience_summary = result.get("audience_summary")
    row.ai_analyzed_at = datetime.utcnow()

    follower_estimate = result.get("follower_estimate")
    if not row.follower_count and isinstance(follower_estimate, (int, float)):
        row.follower_count = int(follower_estimate)

    await db.commit()
    await db.refresh(row)
    return _serialize(row)


# ── GET /summary ─────────────────────────────────────────────────────────────

@router.get("/summary")
async def get_summary(
    months: int = Query(default=12, ge=1, le=60),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """채널별 / 타겟별 / 월별 시딩 비용 집계."""
    today = date.today()
    y, m = today.year, today.month
    for _ in range(months - 1):
        m -= 1
        if m == 0:
            m = 12
            y -= 1
    range_start = date(y, m, 1)

    rows = (
        await db.execute(
            select(InfluencerSeeding).where(InfluencerSeeding.seeded_at >= range_start)
        )
    ).scalars().all()

    by_channel: dict[str, dict] = defaultdict(lambda: {"total_cost": 0.0, "count": 0})
    by_segment: dict[str, dict] = defaultdict(lambda: {"total_cost": 0.0, "count": 0})
    by_month: dict[str, dict] = defaultdict(lambda: {"total_cost": 0.0, "count": 0})

    total_cost = 0.0
    total_count = 0
    analyzed_count = 0

    for r in rows:
        cost = float(r.cost or 0)
        total_cost += cost
        total_count += 1
        if r.ai_target_segment:
            analyzed_count += 1

        by_channel[r.channel]["total_cost"] += cost
        by_channel[r.channel]["count"] += 1

        segment_key = r.ai_target_segment or _UNANALYZED_LABEL
        by_segment[segment_key]["total_cost"] += cost
        by_segment[segment_key]["count"] += 1

        if r.seeded_at:
            month_key = f"{r.seeded_at.year:04d}-{r.seeded_at.month:02d}"
            by_month[month_key]["total_cost"] += cost
            by_month[month_key]["count"] += 1

    return {
        "by_channel": [
            {"channel": k, "total_cost": round(v["total_cost"], 2), "count": v["count"]}
            for k, v in sorted(by_channel.items(), key=lambda kv: -kv[1]["total_cost"])
        ],
        "by_segment": [
            {"segment": k, "total_cost": round(v["total_cost"], 2), "count": v["count"]}
            for k, v in sorted(by_segment.items(), key=lambda kv: -kv[1]["total_cost"])
        ],
        "by_month": [
            {"month": k, "total_cost": round(v["total_cost"], 2), "count": v["count"]}
            for k, v in sorted(by_month.items())
        ],
        "total": {
            "cost": round(total_cost, 2),
            "count": total_count,
            "analyzed_count": analyzed_count,
        },
    }


# ── GET /export ──────────────────────────────────────────────────────────────

@router.get("/export")
async def export_seedings(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """인플루언서 시딩 목록 + 채널별/타겟별 요약 xlsx 다운로드."""
    from io import BytesIO
    from urllib.parse import quote

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = (
        await db.execute(select(InfluencerSeeding).order_by(InfluencerSeeding.seeded_at.desc()))
    ).scalars().all()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2A2D35", end_color="2A2D35", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    bold = Font(bold=True)

    wb = Workbook()

    # ===== Sheet1: 시딩 목록 =====
    ws1 = wb.active
    ws1.title = "시딩 목록"
    headers1 = ["일자", "이름", "채널", "팔로워", "비용", "제품", "AI 타겟", "URL", "메모"]
    ws1.append(headers1)
    for col in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for r in rows:
        ws1.append([
            r.seeded_at.isoformat() if r.seeded_at else "",
            r.name or "",
            r.channel or "",
            r.follower_count or "",
            r.cost or 0,
            r.product or "",
            r.ai_target_segment or "",
            r.url or "",
            r.notes or "",
        ])

    widths1 = [12, 18, 12, 10, 12, 18, 30, 40, 30]
    for i, w in enumerate(widths1, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ===== Sheet2: 채널별 요약 =====
    ws2 = wb.create_sheet("채널별 요약")
    headers2 = ["채널", "건수", "총비용"]
    ws2.append(headers2)
    for col in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    channel_agg: dict[str, dict] = defaultdict(lambda: {"count": 0, "cost": 0.0})
    for r in rows:
        channel_agg[r.channel]["count"] += 1
        channel_agg[r.channel]["cost"] += float(r.cost or 0)
    for channel, agg in sorted(channel_agg.items(), key=lambda kv: -kv[1]["cost"]):
        ws2.append([channel, agg["count"], agg["cost"]])
    for i, w in enumerate([16, 10, 14], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ===== Sheet3: 타겟별 요약 =====
    ws3 = wb.create_sheet("타겟별 요약")
    headers3 = ["AI 타겟", "건수", "총비용"]
    ws3.append(headers3)
    for col in range(1, len(headers3) + 1):
        cell = ws3.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    segment_agg: dict[str, dict] = defaultdict(lambda: {"count": 0, "cost": 0.0})
    for r in rows:
        key = r.ai_target_segment or _UNANALYZED_LABEL
        segment_agg[key]["count"] += 1
        segment_agg[key]["cost"] += float(r.cost or 0)
    for segment, agg in sorted(segment_agg.items(), key=lambda kv: -kv[1]["cost"]):
        ws3.append([segment, agg["count"], agg["cost"]])
    for i, w in enumerate([36, 10, 14], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    today_str = datetime.utcnow().strftime("%Y%m%d")
    filename = f"인플루언서시딩_{today_str}.xlsx"
    quoted = quote(filename)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quoted}",
        },
    )
