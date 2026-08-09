"""Meta 광고 인사이트 API — 스냅샷 + 온디맨드 하이브리드.

엔드포인트:
  GET  /api/v1/insights/trend?days=30  — DB 기반 트렌드 조회
  POST /api/v1/insights/refresh         — 즉시 수집 실행
  GET  /api/v1/insights/status          — 수집기 상태 조회
  GET  /api/v1/insights/export          — 성과분석 엑셀 다운로드 (trend와 동일 파라미터)
"""
import logging
from collections import defaultdict
from datetime import date, timedelta
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.database import get_db
from app.models.user import User
from app.api.v1.endpoints.auth import get_current_user

logger = logging.getLogger(__name__)
router = APIRouter()


# ── 파생 지표 계산 헬퍼 ──────────────────────────────────────────────────────

def _safe_div(numerator: float, denominator: float) -> float:
    """0 나누기를 0으로 처리하는 안전 나눗셈."""
    if not denominator:
        return 0.0
    return round(numerator / denominator, 4)


def _build_series_row(
    row_date: date,
    spend: float,
    impressions: int,
    clicks: int,
    conversions: float,
    revenue: float,
    end_date: Optional[date] = None,
) -> Dict[str, Any]:
    """날짜 + 원본 지표 → 파생 지표 포함 시리즈 행 반환."""
    roas = _safe_div(revenue, spend)
    cpa = _safe_div(spend, conversions)
    ctr = round(_safe_div(clicks, impressions) * 100, 4)
    cpc = round(_safe_div(spend, clicks), 2)
    return {
        "date": row_date.isoformat(),
        "date_end": (end_date or row_date).isoformat(),
        "spend": round(spend, 2),
        "impressions": impressions,
        "clicks": clicks,
        "conversions": round(conversions, 2),
        "revenue": round(revenue, 2),
        "roas": roas,
        "cpa": round(cpa, 2),
        "cpc": cpc,
        "ctr": ctr,
    }


# ── GET /trend ───────────────────────────────────────────────────────────────

def _resolve_trend_range(
    days: int, since: Optional[str], until: Optional[str]
) -> tuple[date, date]:
    """days 또는 since/until 파라미터로부터 조회 범위(since_date, until_date)를 계산."""
    until_date = date.today()
    if since:
        try:
            since_date = date.fromisoformat(since)
            if until:
                until_date = date.fromisoformat(until)
        except ValueError:
            raise HTTPException(status_code=422, detail="since/until 은 YYYY-MM-DD 형식이어야 합니다.")
        if since_date > until_date:
            raise HTTPException(status_code=422, detail="since 가 until 보다 뒤입니다.")
    else:
        since_date = until_date - timedelta(days=days)
    return since_date, until_date


async def _build_insights_trend(
    db: AsyncSession,
    current_user: User,
    since_date: date,
    until_date: date,
    granularity: str,
) -> Dict[str, Any]:
    """DB에 저장된 campaign 레벨 인사이트를 일별/주별로 집계하여 반환. trend·export 공용.

    account.series — 전체 캠페인 합산 지표
    campaigns      — 캠페인별 지표
    """
    from app.models.meta_insight import MetaInsightDaily
    from app.services.meta_insights_collector import collector_state

    # 공유 Meta 자격증명의 ad_account_id 를 기준으로 조회
    # (현재 유저에게 없으면 공유 유저에서 가져옴 — analytics와 동일 패턴)
    from app.api.v1.endpoints.auth import get_shared_meta_credentials

    meta_user: Optional[User] = (
        current_user if current_user.meta_access_token
        else await get_shared_meta_credentials(db)
    )
    ad_account_id: Optional[str] = None
    if meta_user and meta_user.meta_ad_account_id:
        ad_account_id = meta_user.meta_ad_account_id
        if not ad_account_id.startswith("act_"):
            ad_account_id = f"act_{ad_account_id}"

    # DB 조회 — ad_account_id 필터 (없으면 전체)
    query = select(MetaInsightDaily).where(
        MetaInsightDaily.date >= since_date,
        MetaInsightDaily.date <= until_date,
        MetaInsightDaily.level == "campaign",
    )
    if ad_account_id:
        query = query.where(MetaInsightDaily.ad_account_id == ad_account_id)

    result = await db.execute(query.order_by(MetaInsightDaily.date))
    rows = result.scalars().all()

    def _bucket_key(d: date) -> str:
        """granularity에 따른 집계 키 — weekly면 해당 주 월요일 날짜."""
        if granularity == "weekly":
            return (d - timedelta(days=d.weekday())).isoformat()
        return d.isoformat()

    # ── 버킷별 전체 합산 (account 레벨) ──
    account_agg: Dict[str, Dict[str, float]] = defaultdict(
        lambda: {"spend": 0.0, "impressions": 0, "clicks": 0, "conversions": 0.0, "revenue": 0.0}
    )
    # key: campaign_id → {"name": ..., "dates": {date_str: {...}}}
    campaign_agg: Dict[str, Dict[str, Any]] = {}

    for row in rows:
        date_str = _bucket_key(row.date)

        # 계정 집계
        acc = account_agg[date_str]
        acc["spend"] += row.spend
        acc["impressions"] += row.impressions
        acc["clicks"] += row.clicks
        acc["conversions"] += row.conversions
        acc["revenue"] += row.revenue

        # 캠페인 집계
        cid = row.campaign_id or row.object_id
        if cid not in campaign_agg:
            campaign_agg[cid] = {
                "campaign_id": cid,
                "campaign_name": row.campaign_name or row.object_name or cid,
                "dates": defaultdict(
                    lambda: {"spend": 0.0, "impressions": 0, "clicks": 0, "conversions": 0.0, "revenue": 0.0}
                ),
            }
        cdates = campaign_agg[cid]["dates"][date_str]
        cdates["spend"] += row.spend
        cdates["impressions"] += row.impressions
        cdates["clicks"] += row.clicks
        cdates["conversions"] += row.conversions
        cdates["revenue"] += row.revenue

    # ── 시리즈 빌드 ──
    def _row_end(d: date) -> Optional[date]:
        if granularity != "weekly":
            return None
        return min(d + timedelta(days=6), until_date)

    account_series = [
        _build_series_row(
            date.fromisoformat(d),
            v["spend"], v["impressions"], v["clicks"], v["conversions"], v["revenue"],
            end_date=_row_end(date.fromisoformat(d)),
        )
        for d, v in sorted(account_agg.items())
    ]

    campaigns_out = []
    for cid, cdata in campaign_agg.items():
        series = [
            _build_series_row(
                date.fromisoformat(d),
                v["spend"], v["impressions"], v["clicks"], v["conversions"], v["revenue"],
                end_date=_row_end(date.fromisoformat(d)),
            )
            for d, v in sorted(cdata["dates"].items())
        ]
        campaigns_out.append({
            "campaign_id": cdata["campaign_id"],
            "campaign_name": cdata["campaign_name"],
            "series": series,
        })

    return {
        "as_of": collector_state.get("as_of"),
        "account": {"series": account_series},
        "campaigns": campaigns_out,
    }


@router.get("/trend")
async def get_insights_trend(
    days: int = Query(default=30, ge=1, le=400, description="조회 일수 (since/until 미지정 시)"),
    since: Optional[str] = Query(default=None, description="시작일 YYYY-MM-DD (커스텀 범위)"),
    until: Optional[str] = Query(default=None, description="종료일 YYYY-MM-DD (커스텀 범위)"),
    granularity: str = Query(default="daily", description="daily | weekly (주간은 월요일 시작 주 단위 서버 집계)"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """DB에 저장된 campaign 레벨 인사이트를 일별/주별로 집계하여 반환.

    account.series — 전체 캠페인 합산 지표
    campaigns      — 캠페인별 지표
    """
    if granularity not in ("daily", "weekly"):
        raise HTTPException(status_code=422, detail="granularity 는 daily 또는 weekly 여야 합니다.")
    since_date, until_date = _resolve_trend_range(days, since, until)
    return await _build_insights_trend(db, current_user, since_date, until_date, granularity)


# ── GET /hourly-heatmap ──────────────────────────────────────────────────────

# Meta 시간대 breakdown 캐시 (모듈 레벨, TTL 6시간)
_hourly_heatmap_cache: Dict[Any, Any] = {}

_PURCHASE_PRIORITY = ("omni_purchase", "purchase", "offsite_conversion.fb_pixel_purchase")


def _pick_purchase(items: Optional[List[Dict[str, Any]]]) -> float:
    """Meta가 같은 구매를 purchase/omni_purchase/fb_pixel_purchase로 3중 보고 → 단일 선택."""
    if not items:
        return 0.0
    by_type = {str(i.get("action_type")): float(i.get("value") or 0) for i in items}
    for t in _PURCHASE_PRIORITY:
        if t in by_type:
            return by_type[t]
    return 0.0


@router.get("/hourly-heatmap")
async def get_hourly_heatmap(
    days: int = Query(default=30, ge=1, le=92, description="조회 일수 (오늘 포함, since/until 미지정 시)"),
    since: Optional[str] = Query(default=None, description="시작일 YYYY-MM-DD (커스텀 범위)"),
    until: Optional[str] = Query(default=None, description="종료일 YYYY-MM-DD (커스텀 범위)"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """요일×시간대별 Meta 실집행 지표 히트맵 — 광고비/노출/클릭/구매/매출.

    Meta 인사이트 breakdowns=hourly_stats_aggregated_by_advertiser_time_zone +
    time_increment=1 라이브 호출(6h 캐시). 광고 계정 타임존(KST) 기준.
    주의: Meta는 일예산 자동 페이싱이므로 이 분포는 '설정한 예산 배분'이 아니라
    Meta가 실제 지출한 시간대 분포다.
    """
    import time as _time

    import httpx

    from app.api.v1.endpoints.auth import get_shared_meta_credentials
    from app.core.config import get_settings

    since_date, until_date = _resolve_trend_range(days, since, until)
    if since is None:
        since_date = until_date - timedelta(days=days - 1)
    # hourly breakdown은 일×24행이라 범위를 92일로 제한 (초과 시 최근 92일로 클램프)
    clamped = False
    if (until_date - since_date).days > 91:
        since_date = until_date - timedelta(days=91)
        clamped = True
    range_days = (until_date - since_date).days + 1

    meta_user: Optional[User] = (
        current_user if current_user.meta_access_token
        else await get_shared_meta_credentials(db)
    )
    if not meta_user or not meta_user.meta_access_token or not meta_user.meta_ad_account_id:
        return {"available": False, "reason": "Meta 계정이 연결되지 않았습니다."}

    ad_account_id = meta_user.meta_ad_account_id
    if not ad_account_id.startswith("act_"):
        ad_account_id = f"act_{ad_account_id}"

    cache_key = (ad_account_id, since_date.isoformat(), until_date.isoformat())
    cached = _hourly_heatmap_cache.get(cache_key)
    if cached and (_time.time() - cached[0] < 6 * 3600):
        return cached[1]

    settings = get_settings()
    base = f"{settings.META_GRAPH_API_BASE}/{settings.META_API_VERSION}"
    fields_full = "spend,impressions,clicks,actions,action_values"
    fields_lite = "spend,impressions,clicks"

    def _zero_matrix() -> List[List[float]]:
        return [[0.0] * 24 for _ in range(7)]

    matrices: Dict[str, List[List[float]]] = {
        k: _zero_matrix() for k in ("spend", "impressions", "clicks", "purchases", "revenue")
    }
    actions_available = True

    async def _fetch(fields: str) -> Optional[str]:
        """히트맵 행 누적. 성공 시 None, Meta 에러 시 에러 메시지 반환."""
        params: Dict[str, Any] = {
            "access_token": meta_user.meta_access_token,
            "level": "account",
            "fields": fields,
            "breakdowns": "hourly_stats_aggregated_by_advertiser_time_zone",
            "time_increment": 1,
            "time_range": f'{{"since":"{since_date.isoformat()}","until":"{until_date.isoformat()}"}}',
            "limit": 500,
        }
        async with httpx.AsyncClient(timeout=90) as client:
            url = f"{base}/{ad_account_id}/insights"
            while True:
                resp = await client.get(url, params=params)
                body = resp.json()
                if body.get("error"):
                    return str(body["error"].get("message") or body["error"])
                for row in body.get("data", []):
                    try:
                        weekday = date.fromisoformat(str(row.get("date_start"))).weekday()
                        hour = int(str(row.get("hourly_stats_aggregated_by_advertiser_time_zone", "0"))[:2])
                    except (ValueError, TypeError):
                        continue
                    if not (0 <= hour <= 23):
                        continue
                    matrices["spend"][weekday][hour] += float(row.get("spend") or 0)
                    matrices["impressions"][weekday][hour] += float(row.get("impressions") or 0)
                    matrices["clicks"][weekday][hour] += float(row.get("clicks") or 0)
                    matrices["purchases"][weekday][hour] += _pick_purchase(row.get("actions"))
                    matrices["revenue"][weekday][hour] += _pick_purchase(row.get("action_values"))
                next_url = (body.get("paging") or {}).get("next")
                if not next_url:
                    return None
                url, params = next_url, {}

    try:
        err = await _fetch(fields_full)
        if err:
            # 일부 계정은 hourly breakdown + actions 조합을 거부 → 기본 지표만 재시도
            logger.warning(f"[Insights] hourly+actions 실패({err}) — 기본 지표로 재시도")
            actions_available = False
            for m in matrices.values():
                for r in m:
                    r[:] = [0.0] * 24
            err = await _fetch(fields_lite)
            if err:
                logger.error(f"[Insights] hourly heatmap 호출 실패: {err}")
                return {"available": False, "reason": f"Meta API 오류: {err}"}
    except Exception as e:
        logger.error(f"[Insights] hourly heatmap 예외: {e}")
        return {"available": False, "reason": f"Meta API 호출 실패: {e}"}

    for key in ("spend", "revenue"):
        matrices[key] = [[round(v, 2) for v in r] for r in matrices[key]]
    for key in ("impressions", "clicks", "purchases"):
        matrices[key] = [[int(v) for v in r] for r in matrices[key]]

    spend_m = matrices["spend"]
    result = {
        "available": True,
        "since": since_date.isoformat(),
        "until": until_date.isoformat(),
        "days": range_days,
        "clamped": clamped,
        "matrices": matrices,
        "actions_available": actions_available,
        "totals": {k: round(sum(sum(r) for r in m), 2) for k, m in matrices.items()},
        "weekday_spend": [round(sum(r), 2) for r in spend_m],
        "hour_spend": [round(sum(spend_m[w][h] for w in range(7)), 2) for h in range(24)],
        "basis": {
            "spend": "Meta 광고 계정 타임존(KST) 기준, 해당 요일×시간대에 실제 집행된 광고비 합계. "
                     "Meta는 일예산을 자동 페이싱하므로 '설정한 예산 배분'이 아니라 실지출 분포입니다.",
            "attribution": "구매·매출은 노출 발생 시간대 기준 귀속(omni_purchase 우선 단일 선택, 3중 보고 합산 제외).",
            "period": f"{since_date.isoformat()} ~ {until_date.isoformat()} ({range_days}일) 합산. 6시간 캐시."
                      + (" 시간대 분석은 최대 92일까지 지원되어 최근 92일로 잘렸습니다." if clamped else ""),
        },
    }
    _hourly_heatmap_cache[cache_key] = (_time.time(), result)
    return result


# ── POST /refresh ────────────────────────────────────────────────────────────

@router.post("/refresh")
async def refresh_insights(
    since: Optional[str] = Query(default=None, description="YYYY-MM-DD — 지정 시 해당일부터 백필"),
    until: Optional[str] = Query(default=None, description="YYYY-MM-DD — 백필 종료일 (기본 오늘)"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """즉시 Meta 인사이트 수집을 실행하고 결과를 반환. since/until 지정 시 과거 범위 백필."""
    from app.services.meta_insights_collector import collect_insights, collector_state

    for label, v in (("since", since), ("until", until)):
        if v:
            try:
                date.fromisoformat(v)
            except ValueError:
                raise HTTPException(status_code=422, detail=f"{label} 는 YYYY-MM-DD 형식이어야 합니다.")

    try:
        collected_rows = await collect_insights(db, since=since, until=until)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"수집 중 오류 발생: {exc}")

    return {
        "collected_rows": collected_rows,
        "as_of": collector_state.get("as_of"),
    }


# ── GET /status ──────────────────────────────────────────────────────────────

@router.get("/status")
async def get_insights_status(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """수집기 상태 및 DB 요약 반환."""
    from app.models.meta_insight import MetaInsightDaily
    from app.services.meta_insights_collector import collector_state
    from app.services.meta_token_service import token_state

    count_result = await db.execute(
        select(func.count()).select_from(MetaInsightDaily)
    )
    total_rows = count_result.scalar() or 0

    return {
        "as_of": collector_state.get("as_of"),
        "total_rows": total_rows,
        "token_expired": collector_state.get("token_expired", False),
        "last_error": collector_state.get("last_error"),
        "token_expires_at": token_state.get("expires_at"),
        "token_last_refresh": token_state.get("last_refresh"),
    }


# ── GET /export ─────────────────────────────────────────────────────────────

@router.get("/export")
async def export_insights_excel(
    days: int = Query(default=30, ge=1, le=400, description="조회 일수 (since/until 미지정 시)"),
    since: Optional[str] = Query(default=None, description="시작일 YYYY-MM-DD (커스텀 범위)"),
    until: Optional[str] = Query(default=None, description="종료일 YYYY-MM-DD (커스텀 범위)"),
    granularity: str = Query(default="daily", description="daily | weekly (주간은 월요일 시작 주 단위 서버 집계)"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Meta 광고 성과 트렌드를 엑셀(xlsx)로 다운로드. 파라미터는 /trend와 동일."""
    if granularity not in ("daily", "weekly"):
        raise HTTPException(status_code=422, detail="granularity 는 daily 또는 weekly 여야 합니다.")
    since_date, until_date = _resolve_trend_range(days, since, until)
    data = await _build_insights_trend(db, current_user, since_date, until_date, granularity)

    from io import BytesIO
    from urllib.parse import quote

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from fastapi.responses import StreamingResponse

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    money_fmt = "#,##0"
    ratio_fmt = "0.00"

    wb = Workbook()

    # ── Sheet1: 계정 성과 ──
    ws1 = wb.active
    ws1.title = "계정 성과"
    headers1 = ["날짜", "지출", "노출", "클릭", "CTR(%)", "CPC", "전환수", "전환매출", "ROAS", "CPA"]
    for ci, h in enumerate(headers1, start=1):
        cell = ws1.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    money_cols1 = {2, 6, 8, 10}   # 지출/CPC/전환매출/CPA
    ratio_cols1 = {5, 7, 9}       # CTR(%)/전환수/ROAS
    int_cols1 = {3, 4}            # 노출/클릭

    for ri, row in enumerate(data["account"]["series"], start=2):
        date_label = row["date"] if row["date"] == row.get("date_end") else f"{row['date']} ~ {row.get('date_end')}"
        row_vals = [
            date_label, row["spend"], row["impressions"], row["clicks"],
            row["ctr"], row["cpc"], row["conversions"], row["revenue"],
            row["roas"], row["cpa"],
        ]
        for ci, val in enumerate(row_vals, start=1):
            cell = ws1.cell(row=ri, column=ci, value=val)
            if isinstance(val, (int, float)):
                if ci in money_cols1:
                    cell.number_format = money_fmt
                elif ci in ratio_cols1 or ci in int_cols1:
                    cell.number_format = ratio_fmt if ci in ratio_cols1 else "#,##0"

    widths1 = [22, 14, 12, 10, 10, 10, 10, 14, 10, 12]
    for i, w in enumerate(widths1, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = "A2"

    # ── Sheet2: 캠페인 요약 (기간 전체 합산, 지출 내림차순) ──
    ws2 = wb.create_sheet("캠페인 요약")
    headers2 = ["캠페인명", "지출", "전환수", "전환매출", "ROAS"]
    for ci, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    campaign_totals = []
    for camp in data["campaigns"]:
        spend = sum(s["spend"] for s in camp["series"])
        conversions = sum(s["conversions"] for s in camp["series"])
        revenue = sum(s["revenue"] for s in camp["series"])
        roas = round(revenue / spend, 4) if spend else 0.0
        campaign_totals.append({
            "campaign_name": camp["campaign_name"],
            "spend": round(spend, 2),
            "conversions": round(conversions, 2),
            "revenue": round(revenue, 2),
            "roas": roas,
        })
    campaign_totals.sort(key=lambda c: c["spend"], reverse=True)

    for ri, c in enumerate(campaign_totals, start=2):
        ws2.cell(row=ri, column=1, value=c["campaign_name"])
        ws2.cell(row=ri, column=2, value=c["spend"]).number_format = money_fmt
        ws2.cell(row=ri, column=3, value=c["conversions"]).number_format = ratio_fmt
        ws2.cell(row=ri, column=4, value=c["revenue"]).number_format = money_fmt
        ws2.cell(row=ri, column=5, value=c["roas"]).number_format = ratio_fmt

    widths2 = [30, 14, 10, 14, 10]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"성과분석_{since_date.isoformat()}_{until_date.isoformat()}.xlsx"
    quoted = quote(filename)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quoted}"},
    )
