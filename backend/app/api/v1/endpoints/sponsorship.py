"""협찬(스폰서십) 관리 모듈 — 대학축제/동아리/마라톤/학회 등 행사 제품 협찬 기록·집계.

엔드포인트:
  GET    /api/v1/sponsorship/events        — 협찬 목록
  POST   /api/v1/sponsorship/events        — 협찬 등록
  PUT    /api/v1/sponsorship/events/{id}   — 협찬 수정
  DELETE /api/v1/sponsorship/events/{id}   — 협찬 삭제
  GET    /api/v1/sponsorship/summary       — 월별/제품별/행사유형별/조건별 집계
  GET    /api/v1/sponsorship/export        — xlsx 다운로드
"""
import logging
from collections import defaultdict
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.v1.endpoints.auth import get_current_user
from app.db.database import get_db
from app.models.sponsorship import Sponsorship
from app.models.user import User

logger = logging.getLogger(__name__)
router = APIRouter()

EVENT_TYPE_LABELS = {
    "festival": "대학축제",
    "club": "동아리",
    "marathon": "마라톤",
    "conference": "학회",
    "etc": "기타",
}


# ── 공용 헬퍼 ────────────────────────────────────────────────────────────────

def _serialize(row: Sponsorship) -> dict:
    return {
        "id": row.id,
        "target_name": row.target_name,
        "event_type": row.event_type,
        "sponsored_at": row.sponsored_at.isoformat() if row.sponsored_at else None,
        "product": row.product,
        "quantity": row.quantity,
        "estimated_value": row.estimated_value,
        "reason": row.reason,
        "expected_effect": row.expected_effect,
        "conditions": row.conditions,
        "notes": row.notes,
        "created_at": row.created_at.isoformat() if row.created_at else None,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
    }


def _parse_date(value: str, field_name: str = "sponsored_at") -> date:
    try:
        return date.fromisoformat(value)
    except (ValueError, TypeError):
        raise HTTPException(status_code=422, detail=f"{field_name}는 YYYY-MM-DD 형식이어야 합니다.")


def _split_conditions(conditions: Optional[str]) -> list[str]:
    if not conditions:
        return []
    return [c.strip() for c in conditions.split(",") if c.strip()]


# ── GET /events ──────────────────────────────────────────────────────────────

@router.get("/events")
async def list_events(
    event_type: Optional[str] = Query(default=None),
    since: Optional[str] = Query(default=None, description="YYYY-MM-DD"),
    until: Optional[str] = Query(default=None, description="YYYY-MM-DD"),
    q: Optional[str] = Query(default=None, description="대상명/제품/사유 검색어"),
    limit: int = Query(default=300, ge=1, le=1000),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """협찬 목록 (sponsored_at desc)."""
    query = select(Sponsorship)
    if event_type:
        query = query.where(Sponsorship.event_type == event_type)
    if since:
        query = query.where(Sponsorship.sponsored_at >= _parse_date(since, "since"))
    if until:
        query = query.where(Sponsorship.sponsored_at <= _parse_date(until, "until"))
    if q:
        like = f"%{q.strip()}%"
        query = query.where(
            (Sponsorship.target_name.ilike(like))
            | (Sponsorship.product.ilike(like))
            | (Sponsorship.reason.ilike(like))
        )
    query = query.order_by(Sponsorship.sponsored_at.desc(), Sponsorship.id.desc()).limit(limit)

    rows = (await db.execute(query)).scalars().all()
    return {"events": [_serialize(r) for r in rows], "count": len(rows)}


# ── POST /events ─────────────────────────────────────────────────────────────

class SponsorshipCreate(BaseModel):
    target_name: str
    event_type: str
    sponsored_at: str  # YYYY-MM-DD
    product: str
    quantity: Optional[int] = 0
    estimated_value: Optional[float] = None
    reason: Optional[str] = None
    expected_effect: Optional[str] = None
    conditions: Optional[str] = None
    notes: Optional[str] = None


@router.post("/events")
async def create_event(
    payload: SponsorshipCreate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """협찬 등록."""
    if not payload.target_name or not payload.target_name.strip():
        raise HTTPException(status_code=422, detail="target_name은 필수입니다.")
    if not payload.event_type or not payload.event_type.strip():
        raise HTTPException(status_code=422, detail="event_type은 필수입니다.")
    if not payload.product or not payload.product.strip():
        raise HTTPException(status_code=422, detail="product는 필수입니다.")

    sponsored_at = _parse_date(payload.sponsored_at)

    row = Sponsorship(
        target_name=payload.target_name.strip(),
        event_type=payload.event_type.strip(),
        sponsored_at=sponsored_at,
        product=payload.product.strip(),
        quantity=payload.quantity or 0,
        estimated_value=payload.estimated_value,
        reason=payload.reason,
        expected_effect=payload.expected_effect,
        conditions=payload.conditions,
        notes=payload.notes,
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return _serialize(row)


# ── PUT /events/{id} ─────────────────────────────────────────────────────────

class SponsorshipUpdate(BaseModel):
    target_name: Optional[str] = None
    event_type: Optional[str] = None
    sponsored_at: Optional[str] = None
    product: Optional[str] = None
    quantity: Optional[int] = None
    estimated_value: Optional[float] = None
    reason: Optional[str] = None
    expected_effect: Optional[str] = None
    conditions: Optional[str] = None
    notes: Optional[str] = None


@router.put("/events/{event_id}")
async def update_event(
    event_id: int,
    payload: SponsorshipUpdate,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """협찬 부분 수정 (지정된 필드만 갱신)."""
    row = (
        await db.execute(select(Sponsorship).where(Sponsorship.id == event_id))
    ).scalar_one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="협찬 항목을 찾을 수 없습니다.")

    data = payload.model_dump(exclude_none=True)
    if "sponsored_at" in data:
        data["sponsored_at"] = _parse_date(data["sponsored_at"])

    for field, value in data.items():
        setattr(row, field, value)

    await db.commit()
    await db.refresh(row)
    return _serialize(row)


# ── DELETE /events/{id} ──────────────────────────────────────────────────────

@router.delete("/events/{event_id}")
async def delete_event(
    event_id: int,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    row = (
        await db.execute(select(Sponsorship).where(Sponsorship.id == event_id))
    ).scalar_one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="협찬 항목을 찾을 수 없습니다.")
    await db.delete(row)
    await db.commit()
    return {"status": "deleted", "id": event_id}


# ── GET /summary ─────────────────────────────────────────────────────────────

@router.get("/summary")
async def get_summary(
    months: int = Query(default=12, ge=1, le=60),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """월별 / 제품별 / 행사유형별 / 협찬조건별 집계."""
    today = date.today()
    y, m = today.year, today.month
    for _ in range(months - 1):
        m -= 1
        if m == 0:
            m = 12
            y -= 1
    range_start = date(y, m, 1)

    rows = (
        await db.execute(
            select(Sponsorship).where(Sponsorship.sponsored_at >= range_start)
        )
    ).scalars().all()

    by_month: dict[str, dict] = defaultdict(lambda: {"count": 0, "quantity": 0, "estimated_value": 0.0})
    by_product: dict[str, dict] = defaultdict(lambda: {"count": 0, "quantity": 0})
    by_event_type: dict[str, dict] = defaultdict(lambda: {"count": 0, "quantity": 0, "estimated_value": 0.0})
    by_condition: dict[str, int] = defaultdict(int)

    total_count = 0
    total_quantity = 0
    total_estimated_value = 0.0

    for r in rows:
        quantity = int(r.quantity or 0)
        value = float(r.estimated_value or 0)

        total_count += 1
        total_quantity += quantity
        total_estimated_value += value

        if r.sponsored_at:
            month_key = f"{r.sponsored_at.year:04d}-{r.sponsored_at.month:02d}"
            by_month[month_key]["count"] += 1
            by_month[month_key]["quantity"] += quantity
            by_month[month_key]["estimated_value"] += value

        by_product[r.product]["count"] += 1
        by_product[r.product]["quantity"] += quantity

        by_event_type[r.event_type]["count"] += 1
        by_event_type[r.event_type]["quantity"] += quantity
        by_event_type[r.event_type]["estimated_value"] += value

        for cond in _split_conditions(r.conditions):
            by_condition[cond] += 1

    return {
        "by_month": [
            {
                "month": k,
                "count": v["count"],
                "quantity": v["quantity"],
                "estimated_value": round(v["estimated_value"], 2),
            }
            for k, v in sorted(by_month.items())
        ],
        "by_product": [
            {"product": k, "count": v["count"], "quantity": v["quantity"]}
            for k, v in sorted(by_product.items(), key=lambda kv: -kv[1]["count"])
        ],
        "by_event_type": [
            {
                "event_type": k,
                "count": v["count"],
                "quantity": v["quantity"],
                "estimated_value": round(v["estimated_value"], 2),
            }
            for k, v in sorted(by_event_type.items(), key=lambda kv: -kv[1]["count"])
        ],
        "by_condition": [
            {"condition": k, "count": v}
            for k, v in sorted(by_condition.items(), key=lambda kv: -kv[1])
        ],
        "total": {
            "count": total_count,
            "quantity": total_quantity,
            "estimated_value": round(total_estimated_value, 2),
        },
    }


# ── GET /export ──────────────────────────────────────────────────────────────

@router.get("/export")
async def export_events(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """협찬 목록 + 제품별/행사유형별 요약 xlsx 다운로드."""
    from io import BytesIO
    from urllib.parse import quote

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = (
        await db.execute(select(Sponsorship).order_by(Sponsorship.sponsored_at.desc()))
    ).scalars().all()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2A2D35", end_color="2A2D35", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")

    wb = Workbook()

    # ===== Sheet1: 협찬 목록 =====
    ws1 = wb.active
    ws1.title = "협찬 목록"
    headers1 = ["일자", "대상명", "행사유형", "제품", "수량", "환산금액", "협찬조건", "사유", "기대효과", "메모"]
    ws1.append(headers1)
    for col in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for r in rows:
        ws1.append([
            r.sponsored_at.isoformat() if r.sponsored_at else "",
            r.target_name or "",
            EVENT_TYPE_LABELS.get(r.event_type, r.event_type or ""),
            r.product or "",
            r.quantity or 0,
            r.estimated_value or 0,
            r.conditions or "",
            r.reason or "",
            r.expected_effect or "",
            r.notes or "",
        ])

    widths1 = [12, 26, 12, 18, 8, 12, 30, 30, 30, 24]
    for i, w in enumerate(widths1, start=1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # ===== Sheet2: 제품별 요약 =====
    ws2 = wb.create_sheet("제품별 요약")
    headers2 = ["제품", "건수", "수량", "환산금액"]
    ws2.append(headers2)
    for col in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    product_agg: dict[str, dict] = defaultdict(lambda: {"count": 0, "quantity": 0, "estimated_value": 0.0})
    for r in rows:
        product_agg[r.product]["count"] += 1
        product_agg[r.product]["quantity"] += int(r.quantity or 0)
        product_agg[r.product]["estimated_value"] += float(r.estimated_value or 0)
    for product, agg in sorted(product_agg.items(), key=lambda kv: -kv[1]["count"]):
        ws2.append([product, agg["count"], agg["quantity"], agg["estimated_value"]])
    for i, w in enumerate([24, 10, 10, 14], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ===== Sheet3: 행사유형별 요약 =====
    ws3 = wb.create_sheet("행사유형별 요약")
    headers3 = ["행사유형", "건수", "수량", "환산금액"]
    ws3.append(headers3)
    for col in range(1, len(headers3) + 1):
        cell = ws3.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    event_type_agg: dict[str, dict] = defaultdict(lambda: {"count": 0, "quantity": 0, "estimated_value": 0.0})
    for r in rows:
        event_type_agg[r.event_type]["count"] += 1
        event_type_agg[r.event_type]["quantity"] += int(r.quantity or 0)
        event_type_agg[r.event_type]["estimated_value"] += float(r.estimated_value or 0)
    for event_type, agg in sorted(event_type_agg.items(), key=lambda kv: -kv[1]["count"]):
        label = EVENT_TYPE_LABELS.get(event_type, event_type)
        ws3.append([label, agg["count"], agg["quantity"], agg["estimated_value"]])
    for i, w in enumerate([16, 10, 10, 14], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    today_str = datetime.utcnow().strftime("%Y%m%d")
    filename = f"협찬관리_{today_str}.xlsx"
    quoted = quote(filename)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quoted}",
        },
    )
