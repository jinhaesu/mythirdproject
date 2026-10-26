"""광고비 일보 엑셀 업로드 — 시트 xlsx 파싱 + 시스템 매체 자동 연결.

팀이 쓰는 구글시트 「26년 N월 영업 광고비 일보_관리」를 xlsx로 내려받아 그대로 올리면
그 달 일보가 시트와 같은 숫자가 되도록 한다. 공식 숫자 = 'Marketing Cost _ Total' 행의
합계 칸(그 수식 SUM 범위 안 행들의 합계열). 시트↔시스템 재이관(2026-10) 때 확인한 함정을
그대로 처리한다.

1. 일별 없이 합계 칸에만 금액 → 말일에 일괄 반영
2. 합계 칸이 일별 합과 다름(수기 보정) → 작으면 일별 비례 축소, 크면 차액을 말일에 더함
3. 합계 칸이 빈 행 → 공식 합계에서 빠지므로 금액 제외(Limit만 반영)
4. 'Marketing Cost _ Total' 수식 SUM 범위 밖 행 → 제외
5. 매체명 빈칸 → 같은 행 유입·분류 칸을 '_'로 이어 이름, 둘 다 없으면 바로 위 매체에 합산

일자는 머리글 날짜값이 아니라 '이어진 날짜 칸의 위치'로 매긴다(전년도 시트를 복사해
날짜 일련번호가 1년 전인 파일이 있음). DB는 건드리지 않는 순수 함수만 둔다.
"""
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from functools import lru_cache
from io import BytesIO
from typing import Optional

EXCEL_EPOCH = date(1899, 12, 30)
MAX_ROWS = 400  # 데이터 행 상한 (시트는 100행 안팎)
MAX_UNPACKED = 120 * 1024 * 1024  # 압축 푼 크기 상한 (시트 xlsx는 수 MB)

INFLOW_HEADERS = {"유입", "판매채널", "유입채널", "판매처"}
GROUP_HEADERS = {"광고분류", "소스", "분류2"}
OWNER_HEADERS = {"판매담당자", "담당자"}
_TOTAL_HEADER = re.compile(r"^(sub-?total|total|합계|월합계|총합계|총계)$", re.I)
_TOTAL_ROW = re.compile(r"^(marketing\s*cost.*|(.*\s)?total|합계|총계|총\s*합계)$", re.I)


class XlsxImportError(ValueError):
    """사용자에게 그대로 보여줄 업로드 오류. sheets가 있으면 화면이 시트 고르기를 띄운다."""

    def __init__(self, message: str, sheets: Optional[list] = None):
        super().__init__(message)
        self.sheets = sheets


def norm(v) -> str:
    """공백·줄바꿈을 한 칸으로."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


_LOOSE = re.compile(r"[\s_\-()\[\]{}/·.,+&*~:;'\"!?#]")


@lru_cache(maxsize=16384)
def _loose_s(s: str) -> str:
    return _LOOSE.sub("", norm(s)).lower()


def loose(v) -> str:
    """매칭용 — 공백·기호를 빼고 소문자로 ('11번가 슈팅_퍼모먼스' ≈ '11번가슈팅_퍼모먼스')."""
    return _loose_s(v if isinstance(v, str) else norm(v))


def alias_key(inflow: Optional[str], name: str) -> str:
    return f"{loose(inflow)[:100]}|{loose(name)[:200]}"


def _num(v) -> Optional[float]:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _day_of(v):
    """머리글 칸이 날짜면 (년|None, 월, 일), 아니면 None."""
    d = None
    if isinstance(v, datetime):
        d = v.date()
    elif isinstance(v, date):
        d = v
    elif isinstance(v, (int, float)) and not isinstance(v, bool) and 36526 <= v <= 73050:
        d = EXCEL_EPOCH + timedelta(days=int(v))  # 2000~2099 엑셀 날짜 일련번호
    if d:
        return (d.year, d.month, d.day)
    if isinstance(v, str):  # 시스템 export 머리글 '10.1' 등
        m = re.fullmatch(r"(\d{1,2})\s*[./월]\s*(\d{1,2})\s*일?", v.strip())
        if m and 1 <= int(m.group(1)) <= 12 and 1 <= int(m.group(2)) <= 31:
            return (None, int(m.group(1)), int(m.group(2)))
    return None


@dataclass
class Header:
    row: int
    first_col: int  # 날짜 칸 시작 열 (1부터)
    length: int  # 이어진 날짜 칸 수
    year: Optional[int]
    month: int
    first_day: int


class _Grid:
    """시트 값 2차원 목록 — read_only로 한 번 훑어 담고 ws.cell 대신 v(행, 열)로 읽는다(1부터).

    일반 모드 load_workbook은 탭마다 서식·병합까지 풀어 시트 6개짜리 파일에 3초 넘게 걸린다.
    """

    def __init__(self, ws, max_row: int = 3000, max_col: int = 120):
        self.rows = [list(r) for r in ws.iter_rows(max_row=max_row, max_col=max_col, values_only=True)]
        self.max_row = len(self.rows)
        self.max_column = max((len(r) for r in self.rows), default=0)

    def v(self, r: int, c: int):
        if 1 <= r <= self.max_row:
            row = self.rows[r - 1]
            if 1 <= c <= len(row):
                return row[c - 1]
        return None


def _find_header(ws: _Grid, max_scan: int = 15) -> Optional[Header]:
    """날짜 칸이 가로로 28칸 이상 하루씩 이어진 첫 행."""
    max_col = min(ws.max_column or 0, 80)
    for r in range(1, min(ws.max_row or 0, max_scan) + 1):
        best = None
        run_start, run = None, []
        for c in range(1, max_col + 2):
            d = _day_of(ws.v(r, c)) if c <= max_col else None
            if d and run and d[2] == run[-1][2] + 1:
                run.append(d)
                continue
            if run and len(run) >= 28 and (best is None or len(run) > best[1]):
                best = (run_start, len(run), run[0])
            run_start, run = (c, [d]) if d else (None, [])
        if best:
            c0, n, d0 = best
            return Header(row=r, first_col=c0, length=n, year=d0[0], month=d0[1], first_day=d0[2])
    return None


def _sheet_month(h: Header) -> Optional[str]:
    return f"{h.year:04d}-{h.month:02d}" if h.year else None


@dataclass
class SheetRow:
    key: str  # 'r57' — 같은 파일이면 미리보기·반영이 같은 키
    row: int
    inflow: str
    group: str
    name: str
    owner: str = ""
    amounts: dict = field(default_factory=dict)  # 일(1~31) → 금액(원, 보정 후)
    sheet_total: Optional[float] = None  # 합계 칸 값 (빈칸=None)
    raw_sum: float = 0.0  # 일별 칸 합 (보정 전)
    limit: Optional[float] = None
    note: Optional[str] = None
    flags: list = field(default_factory=list)
    nonnum: list = field(default_factory=list)  # 무시한 글자 칸 ('무상', '미달' …)

    @property
    def total(self) -> float:
        return sum(self.amounts.values())

    @property
    def has_data(self) -> bool:
        return bool(self.amounts) or self.limit is not None


@dataclass
class ParsedSheet:
    sheet: str
    sheets: list  # [{name, month, usable}]
    month: str
    detected_month: Optional[str]
    header_row: int
    total_row: Optional[int]
    official_total: Optional[float]
    official_limit: Optional[float]
    has_note_col: bool
    rows: list
    excluded: list  # 공식 합계에서 빠진 금액 [{row, name, amount, reason}]
    warnings: list


def _pick_sheet(wbv, month: str, sheet: Optional[str]):
    """→ (시트 이름, 머리글, 탭 목록). 머리글은 탭마다 위 15행만 훑어 찾는다."""
    y, m = int(month[:4]), int(month[5:7])
    infos, headers = [], {}
    for ws in wbv.worksheets:
        h = _find_header(_Grid(ws, max_row=15))
        headers[ws.title] = h
        infos.append({"name": ws.title, "month": _sheet_month(h) if h else None,
                      "usable": h is not None})
    if sheet:
        if sheet not in headers:
            # 공백만 다르면 같은 탭 — 구글시트 탭 이름 끝에 공백이 붙은 파일이 많다('광고비 일보 10월 ')
            same = [t for t in headers if norm(t) == norm(sheet)]
            if len(same) != 1:
                raise XlsxImportError(f"'{norm(sheet)}' 시트가 파일에 없습니다", infos)
            sheet = same[0]
        if not headers[sheet]:
            raise XlsxImportError(f"'{sheet}' 시트에서 날짜 머리글(1일~말일이 가로로 이어진 행)을 찾지 못했습니다", infos)
        return sheet, headers[sheet], infos

    def score(name: str, h: Header) -> int:
        s = 0
        if h.month == m and h.year == y:
            s += 4
        elif h.month == m:
            s += 2  # 연도만 다름 (전년 시트 복사 잔재)
        if re.search(rf"(?<!\d){m}\s*월", name):
            s += 2
        if "광고비" in name:
            s += 1
        if "min" in name.lower():
            s -= 1  # 서식 원본 탭
        return s

    cands = [(score(ws.title, headers[ws.title]), i, ws) for i, ws in enumerate(wbv.worksheets) if headers[ws.title]]
    if not cands:
        raise XlsxImportError("광고비 일보 형식을 찾지 못했습니다 — 1일~말일 날짜가 가로로 이어진 머리글 행이 있는 시트가 필요합니다", infos)
    best = max(cands, key=lambda t: (t[0], -t[1]))
    if best[0] < 2:
        raise XlsxImportError(f"{m}월 시트를 찾지 못했습니다 — 시트를 직접 골라 주세요", infos)
    title = best[2].title
    return title, headers[title], infos


def _sum_rows(formula, col_letter: str) -> Optional[set]:
    """'=SUM(AK3:AK94)' → {3..94}. 합계열 참조만 본다. 못 읽으면 None."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    body = " ".join(re.findall(r"SUM\(([^)]*)\)", formula, re.I)) or formula
    rows: set = set()
    for c1, r1, c2, r2 in re.findall(r"\$?([A-Z]{1,3})\$?(\d+)(?:\s*:\s*\$?([A-Z]{1,3})\$?(\d+))?", body):
        if c1 != col_letter or (c2 and c2 != col_letter):
            continue
        a, b = int(r1), int(r2 or r1)
        rows.update(range(min(a, b), max(a, b) + 1))
    return rows or None


def parse_workbook(data: bytes, month: str, sheet: Optional[str] = None) -> ParsedSheet:
    """xlsx 바이트 → 한 달치 행 목록. month='YYYY-MM'(일자 매김 기준)."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            unpacked = sum(i.file_size for i in zf.infolist())
    except zipfile.BadZipFile:
        unpacked = None
    if unpacked is not None and unpacked > MAX_UNPACKED:
        raise XlsxImportError("파일 내용이 너무 큽니다 — 광고비 일보 시트만 담긴 파일인지 확인해 주세요")
    bad_file = "엑셀(.xlsx) 파일을 읽지 못했습니다 — 구글시트에서 파일 > 다운로드 > Microsoft Excel(.xlsx)로 받은 파일을 올려 주세요"
    try:
        wbv = load_workbook(BytesIO(data), read_only=True, data_only=True)  # 계산된 값
        wbf = load_workbook(BytesIO(data), read_only=True, data_only=False)  # 수식 (SUM 범위·계산값 누락 판별)
    except Exception:
        raise XlsxImportError(bad_file)
    try:
        title, h, infos = _pick_sheet(wbv, month, sheet)
        ws, wsf = _Grid(wbv[title]), _Grid(wbf[title])  # read_only는 훑을 때 XML을 읽는다
    except XlsxImportError:
        raise
    except Exception:
        raise XlsxImportError(bad_file)
    finally:
        wbv.close()
        wbf.close()
    y, m = int(month[:4]), int(month[5:7])
    days_in_month = ((date(y + (m == 12), m % 12 + 1, 1)) - timedelta(days=1)).day
    warnings: list = []
    detected = _sheet_month(h)
    if h.month != m:
        warnings.append(f"시트 날짜는 {h.month}월인데 {m}월로 올립니다 — 달을 확인해 주세요")
    elif detected and detected != month:
        warnings.append(f"시트 날짜의 연도({h.year})가 선택한 달({y})과 다릅니다 — 전년 시트를 복사한 파일이면 그대로 진행해도 됩니다")

    hr = h.row
    hdr = lambda c: norm(ws.v(hr, c))  # noqa: E731
    run_end = h.first_col + h.length - 1
    name_col = h.first_col - 1
    if name_col < 1:
        raise XlsxImportError("매체명 열을 찾지 못했습니다 — 날짜 칸 바로 왼쪽이 매체명이어야 합니다")
    label_cols = range(1, name_col)
    inflow_col = next((c for c in label_cols if hdr(c).replace(" ", "") in INFLOW_HEADERS), None)
    group_col = next((c for c in label_cols if hdr(c).replace(" ", "") in GROUP_HEADERS and c != inflow_col), None)
    if group_col is None and name_col - 1 >= 1 and name_col - 1 != inflow_col and hdr(name_col - 1) and hdr(name_col - 1) == hdr(name_col):
        group_col = name_col - 1  # '매체 | 매체' 두 열 = 분류 | 매체명
    owner_col = next((c for c in label_cols if hdr(c).replace(" ", "") in OWNER_HEADERS), None)

    # 날짜 칸 오른쪽: 합계·Limit·비고
    right = range(run_end + 1, min(run_end + 9, (ws.max_column or run_end) + 1))
    total_col = next((c for c in right if _TOTAL_HEADER.match(hdr(c).replace(" ", ""))), None)
    limit_col = next((c for c in right if re.search(r"limit|한도|계획", hdr(c), re.I)), None)
    note_col = next((c for c in right if "비고" in hdr(c)), None)
    if total_col is None and run_end + 1 not in (limit_col, note_col):
        total_col = run_end + 1  # 10월 포맷: 합계 머리글이 빈칸
    if total_col is None:
        raise XlsxImportError("합계 열을 찾지 못했습니다 — 날짜 칸 바로 오른쪽이 월 합계여야 합니다")
    # 날짜 칸과 합계 사이 빈 머리글 칸(28~30일 달의 남는 칸)은 말일 뒤 칸으로 본다
    day_cols = [(h.first_day - 1 + i + 1, h.first_col + i) for i in range(h.length)]
    day_cols += [(None, c) for c in range(run_end + 1, total_col) if not hdr(c)]
    total_letter = get_column_letter(total_col)

    # 'Marketing Cost _ Total' 행 → 공식 합계·SUM 범위
    total_row = None
    for r in range(hr + 1, (ws.max_row or hr) + 1):
        if any(_TOTAL_ROW.search(norm(ws.v(r, c))) for c in range(1, name_col + 1)):
            total_row = r
            break
    end_row = total_row or (ws.max_row or hr) + 1
    in_range = _sum_rows(wsf.v(total_row, total_col), total_letter) if total_row else None
    official_total = _num(ws.v(total_row, total_col)) if total_row else None
    official_limit = _num(ws.v(total_row, limit_col)) if total_row and limit_col else None

    rows: list = []
    excluded: list = []
    missing_cache = 0
    cur_inflow = cur_group = cur_owner = ""
    for r in range(hr + 1, end_row):
        own_inflow = norm(ws.v(r, inflow_col)) if inflow_col else ""
        own_group = norm(ws.v(r, group_col)) if group_col else ""
        own_name = norm(ws.v(r, name_col))
        if own_inflow:
            cur_inflow, cur_group = own_inflow, ""
        if owner_col and norm(ws.v(r, owner_col)):
            cur_owner = norm(ws.v(r, owner_col))  # 담당자 칸은 병합돼 첫 행에만 값

        raw: dict = {}
        nonnum: list = []
        overflow = False
        for day, c in day_cols:
            v = ws.v(r, c)
            n = _num(v)
            if n is None:
                if isinstance(v, str) and v.strip():
                    nonnum.append(v.strip())
                continue
            if not n:
                continue
            if day is None or day > days_in_month:
                day, overflow = days_in_month, True
            raw[day] = raw.get(day, 0.0) + n
        tv = ws.v(r, total_col)
        sheet_total = _num(tv)
        if sheet_total is None and isinstance(wsf.v(r, total_col), str) and str(wsf.v(r, total_col)).startswith("="):
            missing_cache += 1
        limit = _num(ws.v(r, limit_col)) if limit_col else None
        note = norm(ws.v(r, note_col)) or None if note_col else None
        raw_sum = sum(raw.values())
        has_any = bool(raw) or bool(sheet_total) or limit is not None

        if in_range is not None and r not in in_range:
            amt = sheet_total if sheet_total else raw_sum
            if amt:
                excluded.append({"row": r, "name": own_name or "_".join(p for p in (own_inflow, own_group) if p),
                                 "amount": amt, "reason": "공식 합계(Marketing Cost _ Total) 수식 범위 밖"})
            continue
        if not own_name and not has_any:
            if own_group:
                cur_group = own_group
            continue

        flags: list = []
        name, group = own_name, ""
        if own_name:
            if own_group:
                cur_group = own_group
            group = cur_group
        else:
            parts = [p for p in (own_inflow, own_group) if p]
            if parts:
                name = "_".join(parts)
                flags.append("derived_name")

        # 합계 칸 기준 보정 (공식 숫자 = 합계열)
        last = days_in_month
        if sheet_total is None:
            amounts = {}
            if raw_sum:
                flags.append("blank_total")
                excluded.append({"row": r, "name": name or "(이름 없음)", "amount": raw_sum,
                                 "reason": "합계 칸이 비어 공식 합계에서 빠짐"})
        elif abs(sheet_total - raw_sum) < 0.5:
            amounts = raw
        elif not raw_sum:
            amounts = {last: sheet_total}
            flags.append("total_only")
        elif 0 < sheet_total < raw_sum or (sheet_total == 0 and raw_sum > 0):
            ratio = sheet_total / raw_sum
            amounts = {d: v * ratio for d, v in raw.items() if v * ratio}
            flags.append("scaled")
        else:
            amounts = dict(raw)
            amounts[last] = amounts.get(last, 0.0) + (sheet_total - raw_sum)
            flags.append("remainder")
        if overflow:
            flags.append("overflow")
        if any(v < 0 for v in amounts.values()):
            flags.append("negative")

        if not name:  # 이름도 유입·분류도 없는 행 → 바로 위 매체에 합산
            if rows and (amounts or limit is not None):
                prev = rows[-1]
                for d, v in amounts.items():
                    prev.amounts[d] = prev.amounts.get(d, 0.0) + v
                prev.raw_sum += raw_sum
                if sheet_total is not None:
                    prev.sheet_total = (prev.sheet_total or 0.0) + sheet_total
                if limit is not None:
                    prev.limit = (prev.limit or 0.0) + limit
                if "merged" not in prev.flags:
                    prev.flags.append("merged")
            elif amounts:
                excluded.append({"row": r, "name": "(이름 없음)", "amount": sum(amounts.values()),
                                 "reason": "매체명·유입·분류가 모두 빈 첫 행"})
            continue

        rows.append(SheetRow(
            key=f"r{r}", row=r, inflow=cur_inflow, group=group, name=name[:200], owner=cur_owner,
            amounts=amounts, sheet_total=sheet_total, raw_sum=raw_sum, limit=limit, note=note,
            flags=flags, nonnum=nonnum[:6],
        ))
        if len(rows) > MAX_ROWS:
            raise XlsxImportError(f"행이 너무 많습니다 ({MAX_ROWS}행까지) — 광고비 일보 시트가 맞는지 확인해 주세요")

    if missing_cache >= 3:
        raise XlsxImportError("수식의 계산값이 저장되지 않은 파일입니다 — 엑셀에서 한 번 열어 저장하거나, 구글시트에서 xlsx로 다시 내려받아 올려 주세요")
    if not rows:
        raise XlsxImportError("가져올 행이 없습니다 — 매체명과 금액이 있는 행이 필요합니다")
    if total_row is None:
        warnings.append("'Marketing Cost _ Total' 합계 행을 찾지 못해, 머리글 아래 모든 행을 가져옵니다")
    computed = sum(r.total for r in rows)
    if official_total is not None and abs(computed - official_total) >= 1:
        warnings.append(f"가져온 합계 {computed:,.0f}원이 파일 공식 합계 {official_total:,.0f}원과 {computed - official_total:+,.0f}원 다릅니다 — 합계 수식을 확인해 주세요")

    return ParsedSheet(
        sheet=title, sheets=infos, month=month, detected_month=detected, header_row=hr,
        total_row=total_row, official_total=official_total, official_limit=official_limit,
        has_note_col=note_col is not None, rows=rows, excluded=excluded, warnings=warnings,
    )


# ─── 시스템 매체 자동 연결 ──────────────────────────────────────────────────

@dataclass
class MediaInfo:
    id: int
    name: str
    inflow: Optional[str]
    group: Optional[str]
    active: bool
    auto: bool = False
    month_spend: float = 0.0  # 그 달 집행 (보드 기준)
    month_limit: Optional[float] = None
    aliases: set = field(default_factory=set)

    @property
    def has_month(self) -> bool:
        return bool(self.month_spend) or self.month_limit is not None


@dataclass
class Suggestion:
    action: str  # map | new | skip
    media_id: Optional[int] = None
    reason: str = ""  # alias | exact | exact_other | similar | slot | inactive | new | empty | auto
    score: float = 0.0

    def as_dict(self) -> dict:
        return {"action": self.action, "media_id": self.media_id, "reason": self.reason, "score": round(self.score, 2)}


@lru_cache(maxsize=16384)
def _bigrams(s: str) -> frozenset:
    return frozenset(s[i:i + 2] for i in range(len(s) - 1)) if len(s) > 1 else frozenset({s})


def name_sim(a: str, b: str) -> float:
    la, lb = loose(a), loose(b)
    if not la or not lb:
        return 0.0
    if la == lb:
        return 1.0
    short, long_ = sorted((la, lb), key=len)
    cont = 0.6 + 0.35 * len(short) / len(long_) if len(short) >= 3 and short in long_ else 0.0
    ba, bb = _bigrams(la), _bigrams(lb)
    dice = 2 * len(ba & bb) / (len(ba) + len(bb))
    return max(cont, dice)


def _same(a: Optional[str], b: Optional[str]) -> bool:
    return bool(loose(a)) and loose(a) == loose(b)


def _same_name(row: SheetRow, m: MediaInfo) -> bool:
    ln = loose(m.name)
    return ln == loose(row.name) or (bool(row.inflow) and ln == loose(f"{row.name}({row.inflow})"))


_AUTO_SUFFIX = re.compile(r"\s*\(자동\)\s*$")  # 시스템 export의 자동 수집 매체 표시


def plain_match(row: SheetRow, name: Optional[str], inflow: Optional[str]) -> bool:
    """이름·판매채널이 그대로 같은 연결 — 기억(별칭)할 필요가 없다."""
    return loose(name) == loose(row.name) and loose(inflow) == loose(row.inflow)


def suggest(rows: list, media: list) -> dict:
    """행별 기본 연결 {key: Suggestion} — 기억된 연결 > 이름 일치 > (이 달 금액 있는) 비슷한 이름 >
    숨긴 매체 이름 일치 > 비슷한 이름 > 새 매체. 비슷한 이름으로는 한 매체에 한 행만 붙이고,
    판매채널이 다른 매체는 이름이 거의 같을 때만 붙인다."""
    usable = [m for m in media if not m.auto]
    active = [m for m in usable if m.active]
    by_alias: dict = {}
    for m in usable:
        for k in m.aliases:
            by_alias.setdefault(k, m)
    # 파일 판매채널 표기 → 시스템 판매채널 (이름 일치 짝에서 배움, 예: '카페홍보' → '널담은공간')
    hint: dict = {}

    def same_in(row: SheetRow, m: MediaInfo) -> bool:
        return _same(m.inflow, row.inflow) or loose(m.inflow) in hint.get(loose(row.inflow), ())

    def prefer(cands: list, row: SheetRow) -> MediaInfo:
        return sorted(cands, key=lambda m: (not _same(m.inflow, row.inflow), not m.has_month,
                                            not _same(m.group, row.group), m.id))[0]

    out: dict = {}
    used: set = set()

    def take(row: SheetRow, m: MediaInfo, reason: str, score: float):
        out[row.key] = Suggestion("map", m.id, reason, score)
        used.add(m.id)
        if reason == "exact_other" and row.inflow and m.inflow:
            hint.setdefault(loose(row.inflow), set()).add(loose(m.inflow))

    autos = [m for m in media if m.auto]
    for row in rows:  # 자동 수집 매체(메타·네이버 검색광고) — 시스템이 채우므로 시트 금액은 건너뜀
        ln = loose(_AUTO_SUFFIX.sub("", row.name))
        m = next((m for m in autos if loose(m.name) == ln
                  and (not row.inflow or not m.inflow or _same(m.inflow, row.inflow))), None)
        if m:
            out[row.key] = Suggestion("skip", m.id, "auto", 1.0)
    for row in rows:
        m = by_alias.get(alias_key(row.inflow, row.name))
        if m and row.key not in out:
            take(row, m, "alias", 1.0)
    for row in rows:  # 이름 + 판매채널 일치
        if row.key in out:
            continue
        cands = [m for m in active if _same_name(row, m) and _same(m.inflow, row.inflow)]
        if cands:
            take(row, prefer(cands, row), "exact", 1.0)
    name_count: dict = {}
    for row in rows:
        name_count[loose(row.name)] = name_count.get(loose(row.name), 0) + 1
    # 이름만 일치(판매채널 표기가 다름) — 파일에 그 이름이 한 번뿐일 때만. 같은 이름 매체가
    # 여럿이면 이 달 금액 있는 쪽, 그래도 못 고르면 앞에서 배운 판매채널 짝으로 고른다
    for tie_break in (False, True):
        for row in rows:
            if row.key in out or name_count[loose(row.name)] > 1:
                continue
            ex = [m for m in active if m.id not in used and _same_name(row, m)]
            if len(ex) > 1:
                ex = [m for m in ex if same_in(row, m)] if tie_break else [m for m in ex if m.has_month]
            if len(ex) == 1:
                take(row, ex[0], "exact_other", 0.95)

    # 같은 (유입, 분류) 칸에 파일도 시스템도 하나씩뿐이면 이름이 달라도 같은 매체로 본다
    slot_rows: dict = {}
    for row in rows:
        if row.inflow and row.group:
            slot_rows.setdefault((loose(row.inflow), loose(row.group)), []).append(row.key)
    slot_media: dict = {}
    for m in active:
        if m.inflow and m.group:
            slot_media.setdefault((loose(m.inflow), loose(m.group)), []).append(m.id)

    def is_slot(row: SheetRow, m: MediaInfo) -> bool:
        k = (loose(row.inflow), loose(row.group))
        return bool(row.inflow and row.group) and len(slot_rows.get(k, [])) == 1 and slot_media.get(k) == [m.id]

    def fuzzy(min_sim: float, need_month: bool):
        pairs = []
        for row in rows:
            if row.key in out or not row.amounts and row.limit is None:
                continue
            for m in active:
                if m.id in used or (need_month and not m.has_month):
                    continue
                s = name_sim(row.name, m.name)
                slot = is_slot(row, m)
                inflow_ok = same_in(row, m)
                if not inflow_ok and s < 0.8:
                    continue  # 다른 판매채널 매체는 이름이 거의 같을 때만
                group_ok = not row.group or not m.group or _same(m.group, row.group)
                if need_month and not group_ok and s < 0.7:
                    continue  # 이 달 금액이 있어도 분류까지 다르면 이름이 꽤 같아야
                if s < min_sim and not slot:
                    continue
                score = s + 0.15 * inflow_ok + 0.1 * _same(m.group, row.group) + 0.2 * m.has_month
                pairs.append((score, s, slot, row, m))
        for score, s, slot, row, m in sorted(pairs, key=lambda t: (-t[0], t[3].row, t[4].id)):
            if row.key in out or m.id in used:
                continue
            take(row, m, "similar" if s >= min_sim else "slot", min(score, 0.99))

    fuzzy(0.45, need_month=True)
    # 숨긴(비활성) 매체 이름 일치 → 다시 켜서 사용 (판매채널까지 같은 짝 먼저)
    pairs = [(not _same(m.inflow, row.inflow), row.row, m.id, row, m)
             for row in rows if row.key not in out and row.has_data
             for m in usable if not m.active and _same_name(row, m)]
    for *_, row, m in sorted(pairs, key=lambda t: t[:3]):
        if row.key in out or m.id in used:
            continue
        take(row, m, "inactive", 0.9)
    fuzzy(0.6, need_month=False)
    for row in rows:
        if row.key not in out:
            out[row.key] = Suggestion("new", None, "new") if row.has_data else Suggestion("skip", None, "empty")
    return out


def candidates(row: SheetRow, media: list, limit: int = 6) -> list:
    """행별 추천 후보 (화면의 매체 고르기 맨 위)."""
    scored = []
    for m in media:
        if m.auto:
            continue
        s = name_sim(row.name, m.name) + 0.15 * _same(m.inflow, row.inflow) + 0.1 * _same(m.group, row.group)
        if s >= 0.35:
            scored.append((s + 0.05 * m.active, m))
    scored.sort(key=lambda t: (-t[0], t[1].id))
    return [{"media_id": m.id, "score": round(s, 2)} for s, m in scored[:limit]]
