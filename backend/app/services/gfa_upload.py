"""네이버 GFA 성과 보고서 파일(CSV·XLSX) 업로드 — 파일만 읽는 순수 함수(DB는 건드리지 않는다).

GFA 광고 관리자 보고서에서 내려받은 '일별' 성과 파일을 그대로 올리면 성과 대시보드가
그 값으로 그려진다(2026-10-27 마케팅 연다은 건의). 실제 내보내기 견본 없이 만들어서
머리글은 넓게 받고, 못 알아본 칸은 미리보기 화면에서 직접 고르게 한다.

- 머리글 행: 위 40행 중 아는 칸 이름이 가장 많은 행 (날짜·비용 칸을 못 찾으면 직접 고르기)
- 칸 이름: 한글·영문 동의어, 공백·기호 무시. 비용이 VAT 포함/제외 둘 다면 제외 쪽,
  전환은 구매(완료) 칸이 있으면 그쪽, 직접·간접만 있으면 둘을 더한다
- 숫자: 쉼표·원·%·₩ 빼고 읽고 '-'·빈칸은 0
- 날짜: 2026-10-01 · 2026.10.01. · 2026/10/01 · 20261001 · 2026년 10월 1일 · 엑셀 날짜 일련번호,
  시작=끝인 기간('2026.10.01. ~ 2026.10.01.'). 여러 날을 묶은 행은 일별이 아니라서 오류
- 합계·소계 행은 건너뛰고, 같은 (계정·날짜·캠페인·광고그룹·소재) 행은 더한다
"""
import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from functools import lru_cache
from html.parser import HTMLParser
from io import BytesIO
from typing import Optional

EXCEL_EPOCH = date(1899, 12, 30)
MAX_DATA_ROWS = 120_000  # 데이터 행 상한 (소재×일 90일치도 수만 행, 서버 메모리 1Gi)
MAX_COLS = 150
MAX_UNPACKED = 200 * 1024 * 1024  # 압축 푼 크기 상한
HEADER_SCAN = 40  # 머리글 찾는 위쪽 행 수

# 화면에 보이는 순서 그대로
FIELDS = {
    "date": "날짜",
    "account": "광고계정",
    "campaign_name": "캠페인",
    "campaign_id": "캠페인 ID",
    "adgroup_name": "광고그룹",
    "adgroup_id": "광고그룹 ID",
    "creative_name": "소재",
    "creative_id": "소재 ID",
    "creative_type": "소재 유형",
    "product": "상품",
    "promotion": "프로모션",
    "impressions": "노출",
    "clicks": "클릭",
    "spend": "비용",
    "conversions": "전환 수",
    "conv_sales": "전환 매출",
    "budget": "예산",
}
REQUIRED = ("date", "spend")
METRICS = ("impressions", "clicks", "spend", "conversions", "conv_sales")
DIMS = ("campaign_id", "campaign_name", "adgroup_id", "adgroup_name", "creative_id", "creative_name",
        "creative_type", "product", "promotion")
_IGNORE = "_ignore"  # 아는 칸이지만 쓰지 않음(CTR·CPC·ROAS 등은 다시 계산)


class GfaUploadError(ValueError):
    """사용자에게 그대로 보여줄 업로드 오류. sheets가 있으면 화면이 시트 고르기를 띄운다."""

    def __init__(self, message: str, sheets: Optional[list] = None):
        super().__init__(message)
        self.sheets = sheets


@lru_cache(maxsize=65536)
def _norm_s(s: str) -> str:
    return " ".join(s.split())


def norm(v) -> str:
    """공백·줄바꿈을 한 칸으로 (같은 캠페인·날짜 글자가 수만 행 되풀이돼 캐시)."""
    if v is None:
        return ""
    if isinstance(v, str):
        return _norm_s(v)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return _norm_s(str(v))


# ─── 칸 이름 알아보기 ─────────────────────────────────────────────────────────

_STRIP = re.compile(r"[\s_\-./:·•()\[\]{}<>%*~,'\"#|+=?!;₩\\]+")
_VAT_TOKEN = re.compile(
    r"(?:vat|부가세|부가가치세)(?:포함|제외|별도|미포함|included|include|incl|excluded|exclude|excl)"
    r"|(?:included|include|incl|excluded|exclude|excl)vat"
)
_QUAL = re.compile(r"직접|간접")
_WINDOW = re.compile(r"(?:\d+일|당일)(?:기여|클릭|조회|노출)*$")


def hkey(v) -> str:
    """머리글 비교용 — 공백·기호를 빼고 소문자 ('총 비용(VAT 포함)' → '총비용vat포함')."""
    return _STRIP.sub("", norm(v)).lower()


def vat_of(header) -> Optional[str]:
    k = hkey(header)
    m = _VAT_TOKEN.search(k)
    if not m:
        return None
    t = m.group(0)
    return "incl" if ("포함" in t and "미포함" not in t) or "incl" in t else "excl"


def _syn(field_name: str, pri: int, *keys: str) -> dict:
    return {k: (field_name, pri) for k in keys}


_EXACT: dict = {}
for _d in (
    _syn("date", 0, "일", "일자", "날짜", "date", "day", "일별", "통계일", "통계일자", "집계일", "집계일자",
         "기준일", "기준일자", "보고일", "reportingdate", "statdate", "statdt"),
    _syn("date", 1, "일시", "reportingstarts"),
    _syn("date", 2, "기간", "보고기간", "조회기간", "period"),
    _syn("account", 0, "광고계정", "광고계정명", "광고계정이름", "account", "accountname", "adaccount",
         "adaccountname"),
    _syn("account", 1, "계정", "계정명", "계정이름", "광고주", "광고주명"),
    _syn("campaign_name", 0, "캠페인", "캠페인명", "캠페인이름", "campaign", "campaignname"),
    _syn("campaign_id", 0, "캠페인id", "캠페인번호", "캠페인no", "campaignid", "campaignno"),
    _syn("adgroup_name", 0, "광고그룹", "광고그룹명", "광고그룹이름", "adgroup", "adgroupname", "adset",
         "adsetname", "광고세트", "광고세트명", "광고세트이름"),
    _syn("adgroup_id", 0, "광고그룹id", "광고그룹번호", "광고그룹no", "adgroupid", "adsetid", "adsetno",
         "광고세트id"),
    _syn("creative_name", 0, "광고소재", "광고소재명", "광고소재이름", "소재", "소재명", "소재이름",
         "creative", "creativename"),
    _syn("creative_name", 1, "광고", "광고명", "광고이름", "ad", "adname"),
    _syn("creative_id", 0, "광고소재id", "광고소재번호", "소재id", "소재번호", "creativeid", "creativeno"),
    _syn("creative_id", 1, "광고id", "adid"),
    _syn("creative_type", 0, "소재유형", "광고소재유형", "소재형식", "광고소재형식", "소재타입", "소재형태",
         "광고소재형태", "creativetype", "creativeformat"),
    _syn("creative_type", 1, "광고형식", "광고유형", "adformat", "adtype", "format"),
    _syn("creative_type", 2, "형식", "유형", "타입", "type"),
    _syn("product", 0, "상품", "상품명", "상품이름", "제품", "제품명", "product", "productname"),
    _syn("product", 1, "품목"),
    _syn("promotion", 0, "프로모션", "프로모션명", "프로모션이름", "promotion", "promotionname"),
    _syn("promotion", 1, "이벤트", "이벤트명", "기획전"),
    _syn("impressions", 0, "노출", "노출수", "노출횟수", "impressions", "impression", "imps", "imp",
         "impcount"),
    _syn("clicks", 0, "클릭", "클릭수", "클릭횟수", "clicks", "click", "clickcount"),
    _syn("spend", 0, "총비용", "비용", "광고비", "총광고비", "광고비용", "집행금액", "광고집행금액", "집행비용",
         "소진금액", "소진액", "소진비용", "spend", "amountspent", "cost", "totalcost", "adcost"),
    _syn("spend", 1, "사용금액", "지출", "지출금액"),
    # 전환: 구매(완료) 칸이 일반 전환 칸보다 앞선다
    _syn("conversions", 0, "구매완료전환수", "구매완료수", "구매수", "구매건수", "구매전환수", "purchases"),
    _syn("conversions", 1, "구매완료", "구매완료전환", "구매전환", "purchase"),
    _syn("conversions", 2, "구매"),
    _syn("conversions", 10, "전환수", "총전환수", "전환건수", "conversions"),
    _syn("conversions", 11, "전환", "conversion"),
    _syn("conv_sales", 0, "구매완료전환매출", "구매완료전환매출액", "구매완료매출", "구매완료매출액",
         "구매완료전환금액", "구매완료전환값", "구매전환매출", "구매전환매출액", "purchasevalue",
         "purchaseconversionvalue"),
    _syn("conv_sales", 1, "구매전환값", "구매금액", "구매액"),
    _syn("conv_sales", 10, "전환매출", "전환매출액", "총전환매출", "총전환매출액", "전환금액", "전환매출합계",
         "conversionvalue", "conversionrevenue"),
    _syn("conv_sales", 11, "전환값", "전환가치"),
    _syn("conv_sales", 12, "매출", "매출액", "revenue"),
    _syn("budget", 0, "일예산", "일일예산", "하루예산", "캠페인일예산", "광고그룹일예산", "dailybudget"),
    _syn("budget", 1, "예산", "캠페인예산", "광고그룹예산", "budget"),
    _syn("budget", 2, "총예산", "기간예산", "월예산", "totalbudget", "lifetimebudget"),
    _syn(_IGNORE, 0, "ctr", "클릭률", "클릭율", "cpc", "클릭당비용", "평균클릭비용", "평균cpc", "cpm",
         "노출당비용", "1000회노출당비용", "천회노출당비용", "roas", "광고수익률", "전환율", "cvr", "cpa",
         "전환당비용", "구매당비용", "결과당비용", "예산소진율", "소진율", "도달", "도달수", "reach", "빈도",
         "frequency", "동영상조회수", "동영상재생수", "재생수", "조회수", "결과", "결과수", "상태",
         "캠페인상태", "광고그룹상태", "소재상태", "운영상태", "목적", "캠페인목적", "입찰가", "입찰방식",
         "과금방식", "게재위치", "디바이스", "기기", "성별", "연령", "연령대", "지역", "참여수", "좋아요"),
):
    _EXACT.update(_d)

# 행을 잘게 나누는 칸 — 있으면 합쳐서 본다고 알려 준다
_BREAKDOWN = {"게재위치", "디바이스", "기기", "성별", "연령", "연령대", "지역"}


def _fuzzy(base: str) -> Optional[tuple]:
    """정확히 맞는 이름이 없을 때 — 들어 있는 낱말로 짐작 (우선순위 20 이상)."""
    if re.search(r"cp[cmav]|ctr|roas|당비용|률|율|순위|평균", base):
        return (_IGNORE, 20)
    if "예산" in base or "budget" in base:
        return ("budget", 25)
    buy = "구매" in base or "purchase" in base
    if re.search(r"매출|전환금액|전환값|conversionvalue|revenue|purchasevalue", base):
        return ("conv_sales", 20 if buy else 30)
    if re.search(r"전환|구매|conversion|purchase", base):
        return ("conversions", 20 if buy else 30)
    if re.search(r"날짜|일자|기간|date|period", base):  # '집행 기간'이 비용으로 잡히지 않게 비용보다 먼저
        return ("date", 25)
    if re.search(r"비용|광고비|spend|cost|소진|집행", base):
        return ("spend", 25)
    if "노출" in base or "impression" in base:
        return ("impressions", 25)
    if "클릭" in base or "click" in base:
        return ("clicks", 25)
    is_id = bool(re.search(r"id$|번호|no$", base))
    if "캠페인" in base or "campaign" in base:
        return ("campaign_id" if is_id else "campaign_name", 25)
    if re.search(r"광고그룹|adgroup|adset|광고세트", base):
        return ("adgroup_id" if is_id else "adgroup_name", 25)
    if "소재" in base or "creative" in base:
        if re.search(r"유형|형식|타입|형태|type|format", base):
            return ("creative_type", 25)
        return ("creative_id" if is_id else "creative_name", 25)
    if re.search(r"상품|제품|product", base):
        return ("product", 25)
    if re.search(r"프로모션|promotion|이벤트", base):
        return ("promotion", 25)
    if "계정" in base or "account" in base:
        return ("account", 25)
    return None


@dataclass
class Col:
    index: int  # 0부터
    header: str
    field: Optional[str] = None  # FIELDS 키 | _IGNORE | None(모름)
    priority: int = 99
    vat: Optional[str] = None  # incl | excl | None
    qual: str = "total"  # total | direct | indirect
    breakdown: bool = False


def classify(index: int, header) -> Col:
    text = norm(header)
    c = Col(index=index, header=text[:120])
    k = hkey(text)
    if not k:
        return c
    c.vat = vat_of(text)
    base = _VAT_TOKEN.sub("", k)
    if "직접" in base:
        c.qual = "direct"
    elif "간접" in base:
        c.qual = "indirect"
    base = _QUAL.sub("", base)
    c.breakdown = base in _BREAKDOWN
    tries = [(base, 0)]
    if base.endswith("원") or base.endswith("krw"):
        tries.append((re.sub(r"(원|krw)$", "", base), 0))
    if base.startswith("총") or base.startswith("전체"):
        tries.append((re.sub(r"^(총|전체)", "", base), 1))
    w = _WINDOW.sub("", base)
    if w != base:
        tries.append((w, 5))
    for t, extra in tries:
        hit = _EXACT.get(t)
        if hit:
            c.field, c.priority = hit[0], hit[1] + extra
            return c
    hit = _fuzzy(base)
    if hit:
        c.field, c.priority = hit
    return c


def col_letter(i: int) -> str:
    """0 → A, 25 → Z, 26 → AA (엑셀 열 이름)."""
    s, n = "", i + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _header_score(row: list) -> tuple:
    cols = [classify(i, v) for i, v in enumerate(row[:MAX_COLS])]
    fields = {c.field for c in cols if c.field and c.field != _IGNORE}
    known = sum(1 for c in cols if c.field)
    score = len(fields) + (2 if "date" in fields else 0) + (2 if "spend" in fields else 0)
    return score, len(fields), known


def find_header(rows: list) -> Optional[int]:
    """아는 칸 이름이 가장 많은 행(0부터). 두 칸도 못 알아보면 None."""
    best, best_i = None, None
    for i, row in enumerate(rows[:HEADER_SCAN]):
        sc = _header_score(list(row))
        if sc[1] < 2:
            continue
        if best is None or sc > best:
            best, best_i = sc, i
    return best_i


@dataclass
class Mapping:
    cols: list  # [Col] 머리글 칸 전부
    fields: dict  # field → [열 index] (전환 직접+간접은 두 칸)
    manual: set  # 사람이 고른 field
    notes: list  # 화면에 보일 안내


def pick_columns(cols: list, manual: Optional[dict] = None) -> Mapping:
    """칸 → 항목. 같은 항목 후보가 여럿이면 우선순위(비용은 VAT 제외 > 표시 없음 > 포함)."""
    notes: list = []
    out: dict = {}
    vat_rank = {"excl": 0, None: 1, "incl": 2}
    qual_rank = {"total": 0, "direct": 1, "indirect": 2}
    for f in FIELDS:
        cands = [c for c in cols if c.field == f]
        if not cands:
            continue
        if f == "spend":
            exact = [c for c in cands if c.priority < 20] or cands
            exact.sort(key=lambda c: (vat_rank[c.vat], c.priority, c.index))
            out[f] = [exact[0].index]
            vats = {c.vat for c in exact}
            if exact[0].vat == "excl" and "incl" in vats:
                inc = next(c for c in exact if c.vat == "incl")
                notes.append(f"비용은 VAT 제외 칸 「{exact[0].header}」을 썼습니다 (VAT 포함 칸 「{inc.header}」은 쓰지 않음)")
            continue
        if f in ("conversions", "conv_sales"):
            cands.sort(key=lambda c: (c.priority, qual_rank[c.qual], c.index))
            top = cands[0]
            if top.qual == "total":
                out[f] = [top.index]
            else:
                same = [c for c in cands if c.priority == top.priority]
                d = next((c for c in same if c.qual == "direct"), None)
                ind = next((c for c in same if c.qual == "indirect"), None)
                if d and ind:
                    out[f] = [d.index, ind.index]
                    notes.append(f"{FIELDS[f]}: 「{d.header}」+「{ind.header}」을 더했습니다")
                else:
                    out[f] = [top.index]
                    notes.append(f"{FIELDS[f]}: 「{top.header}」만 있어 그 값만 씁니다")
            continue
        cands.sort(key=lambda c: (c.priority, c.index))
        out[f] = [cands[0].index]

    chosen: set = set()
    for f, v in (manual or {}).items():
        if f not in FIELDS:
            raise GfaUploadError(f"모르는 항목입니다: {f}")
        if v is None or v == "" or v == -1:
            out.pop(f, None)
            chosen.add(f)
            continue
        idxs = v if isinstance(v, list) else [v]
        try:
            idxs = [int(x) for x in idxs]
        except (TypeError, ValueError):
            raise GfaUploadError(f"{FIELDS[f]} 칸 번호가 올바르지 않습니다")
        for x in idxs:
            if not 0 <= x < len(cols):
                raise GfaUploadError(f"{FIELDS[f]}: 파일에 없는 칸입니다 ({x + 1}번째)")
        out[f] = idxs
        chosen.add(f)

    used: dict = {}
    for f in METRICS + ("budget",):
        for x in out.get(f, []):
            used.setdefault(x, []).append(f)
    for x, fs in used.items():
        if len(fs) > 1:
            notes.append(f"「{cols[x].header or col_letter(x) + '열'}」 칸을 {', '.join(FIELDS[f] for f in fs)}에 함께 썼습니다 — 맞는지 확인해 주세요")
    return Mapping(cols=cols, fields=out, manual=chosen, notes=notes)


# ─── 값 읽기 ─────────────────────────────────────────────────────────────────

_DATE_RE = re.compile(r"(\d{4})\s*[-./년]\s*(\d{1,2})\s*[-./월]\s*(\d{1,2})")
_TOTAL_EXACT = {"합계", "총계", "소계", "총합계", "전체합계", "total", "subtotal", "sub total",
                "sub-total", "grand total", "sum", "합계:", "총계:"}
_TOTAL_RE = re.compile(r"^(합\s*계|총\s*계|소\s*계|total|sub-?\s*total|grand\s*total)\b|(합계|소계|총계)\s*:?$", re.I)
_NUM_JUNK = re.compile(r"[,\s₩원%회건배]|krw", re.I)
_EMPTY_NUM = {"", "-", "–", "—", "n/a", "na", "null", "none", "--"}


@lru_cache(maxsize=65536)
def _total_s(s: str, loose: bool) -> bool:
    return s in _TOTAL_EXACT or bool(_TOTAL_RE.search(s)) or (loose and s in ("전체", "누계"))


def is_total(v, loose: bool = False) -> bool:
    """합계·소계 표시. '전체'·'누계'는 날짜 칸에서만(loose) — 광고그룹 이름이 '전체'인 경우가 있다."""
    s = norm(v).lower()
    return bool(s) and _total_s(s, loose)


def parse_day(v) -> tuple:
    """→ ('ok', date) | ('empty'|'total'|'range'|'bad', None)."""
    if v is None:
        return "empty", None
    if isinstance(v, datetime):
        return "ok", v.date()
    if isinstance(v, date):
        return "ok", v
    if isinstance(v, bool):
        return "bad", None
    if isinstance(v, (int, float)):
        if v != v:  # NaN
            return "empty", None
        if 36526 <= v <= 73050:  # 2000~2099 엑셀 날짜 일련번호
            return "ok", EXCEL_EPOCH + timedelta(days=int(v))
        if float(v).is_integer() and 20000101 <= int(v) <= 20991231:
            iv = int(v)
            try:
                return "ok", date(iv // 10000, iv // 100 % 100, iv % 100)
            except ValueError:
                return "bad", None
        return "bad", None
    return _day_s(norm(v))


@lru_cache(maxsize=8192)
def _day_s(s: str) -> tuple:
    if not s:
        return "empty", None
    if is_total(s, loose=True):
        return "total", None
    found = []
    for m in _DATE_RE.finditer(s):
        try:
            found.append(date(int(m.group(1)), int(m.group(2)), int(m.group(3))))
        except ValueError:
            return "bad", None
    if not found:
        if re.fullmatch(r"\d{8}", s):
            try:
                return "ok", date(int(s[:4]), int(s[4:6]), int(s[6:]))
            except ValueError:
                return "bad", None
        if re.fullmatch(r"\d{5}(\.0+)?", s):
            return parse_day(float(s))
        return "bad", None
    if len(set(found)) > 1:
        return "range", None
    return "ok", found[0]


def parse_num(v) -> tuple:
    """→ (값, 숫자로 읽었는지). '-'·빈칸은 0, '(1,234)'는 음수."""
    if v is None:
        return 0.0, True
    if isinstance(v, bool):
        return 0.0, False
    if isinstance(v, (int, float)):
        return (0.0, True) if v != v else (float(v), True)
    s = str(v).strip()
    try:  # 대부분 '1,234' — 정규식 없이 먼저
        x = float(s.replace(",", ""))
        if x == x and abs(x) != float("inf"):
            return x, True
    except ValueError:
        pass
    if s.lower() in _EMPTY_NUM:
        return 0.0, True
    neg = s.startswith("(") and s.endswith(")")
    s = _NUM_JUNK.sub("", s.strip("()"))
    if s.lower() in _EMPTY_NUM:
        return 0.0, True
    try:
        x = float(s)
    except ValueError:
        return 0.0, False
    if x != x or abs(x) == float("inf"):  # 'nan'·'inf' 글자
        return 0.0, False
    return (-x if neg else x), True


def _text(v, limit: int) -> Optional[str]:
    s = norm(v)
    return s[:limit] if s else None


def type_group(t: Optional[str]) -> str:
    """소재 유형 → 이미지형 / 동영상형 / 기타 (대시보드 비교용)."""
    s = (t or "").lower()
    if not s:
        return "미지정"
    if re.search(r"동영상|영상|video|movie|vod", s):
        return "동영상형"
    if re.search(r"이미지|image|배너|banner|슬라이드|slide|캐러셀|carousel|컬렉션|collection|네이티브|native|카탈로그|catalog", s):
        return "이미지형"
    return "기타"


# ─── 파일 열기 ───────────────────────────────────────────────────────────────

class _TableParser(HTMLParser):
    """엑셀(.xls)로 받았는데 실은 HTML 표인 파일."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows: list = []
        self._row = None
        self._cell: list = []
        self._in = False

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._in, self._cell = True, []
        elif tag == "br" and self._in:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._in:
            self._row.append("".join(self._cell).strip())
            self._in = False
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._in:
            self._cell.append(data)


def _decode(data: bytes) -> tuple:
    if data[:3] == b"\xef\xbb\xbf":
        return data[3:].decode("utf-8", errors="replace"), "utf-8"
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace"), "utf-16"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp949"), "cp949"
    except UnicodeDecodeError:
        return data.decode("cp949", errors="replace"), "cp949"


def _csv_rows(data: bytes) -> tuple:
    text, enc = _decode(data)
    lines = text.splitlines()
    if lines and re.fullmatch(r"sep=.", lines[0].strip(), re.I):  # 엑셀용 'sep=,' 첫 줄
        delim = lines[0].strip()[4]
        text = "\n".join(lines[1:])
    else:
        sample = "\n".join(lines[:30])
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
        except csv.Error:
            counts = {d: sample.count(d) for d in (",", "\t", ";", "|")}
            delim = max(counts, key=counts.get) if any(counts.values()) else ","
    rows = []
    for r in csv.reader(io.StringIO(text), delimiter=delim):
        rows.append(r[:MAX_COLS])
        if len(rows) > MAX_DATA_ROWS + HEADER_SCAN:
            raise GfaUploadError(f"행이 너무 많습니다 (최대 {MAX_DATA_ROWS:,}행) — 기간을 나눠 올려 주세요")
    return rows, enc, {"\t": "탭", ",": "쉼표", ";": "세미콜론", "|": "세로줄"}.get(delim, delim)


def _html_rows(data: bytes) -> tuple:
    text, enc = _decode(data)
    m = re.search(r"charset=[\"']?([\w-]+)", text[:4096], re.I)
    if m and m.group(1).lower() in ("euc-kr", "ks_c_5601-1987", "cp949") and enc == "utf-8":
        text, enc = data.decode("cp949", errors="replace"), "cp949"
    p = _TableParser()
    p.feed(text)
    if len(p.rows) > MAX_DATA_ROWS + HEADER_SCAN:
        raise GfaUploadError(f"행이 너무 많습니다 (최대 {MAX_DATA_ROWS:,}행) — 기간을 나눠 올려 주세요")
    return [r[:MAX_COLS] for r in p.rows], enc


def _xlsx_open(data: bytes):
    try:
        with zipfile.ZipFile(BytesIO(data)) as z:
            unpacked = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile:
        raise GfaUploadError("엑셀 파일을 열 수 없습니다 — 파일이 손상됐거나 엑셀(.xlsx)이 아닙니다")
    if unpacked > MAX_UNPACKED:
        raise GfaUploadError("엑셀 파일 내용이 너무 큽니다 — 기간을 나눠 올려 주세요")
    from openpyxl import load_workbook

    try:
        return load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — 사용자에게 보여줄 메시지로 바꾼다
        raise GfaUploadError(f"엑셀 파일을 읽지 못했습니다 ({type(e).__name__}) — CSV로 내려받아 올려 보세요")


def _ws_rows(ws, max_row: Optional[int] = None, min_row: int = 1, max_col: Optional[int] = None) -> list:
    try:
        ws.reset_dimensions()  # 잘못 적힌 범위(A1 등) 때문에 칸이 잘리지 않게
    except AttributeError:
        pass
    out = []
    for r in ws.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col, values_only=True):
        out.append(list(r[:MAX_COLS]))
        if len(out) > MAX_DATA_ROWS + HEADER_SCAN:
            raise GfaUploadError(f"행이 너무 많습니다 (최대 {MAX_DATA_ROWS:,}행) — 기간을 나눠 올려 주세요")
    return out


@dataclass
class Grid:
    kind: str  # xlsx | csv | html
    rows: list  # 머리글 위쪽 포함 전체 행 (값 목록)
    sheet: Optional[str] = None
    sheets: list = field(default_factory=list)  # [{name, usable}]
    encoding: Optional[str] = None
    delimiter: Optional[str] = None


def read_grid(data: bytes, filename: str = "", sheet: Optional[str] = None) -> Grid:
    """파일 → 행 목록. xlsx는 머리글이 가장 잘 맞는 시트(또는 고른 시트)."""
    head = data[:2048].lower()
    if data[:2] == b"PK":
        wb = _xlsx_open(data)
        try:
            infos, best = [], None
            for ws in wb.worksheets:
                top = _ws_rows(ws, max_row=HEADER_SCAN)
                hi = find_header(top)
                sc = _header_score(top[hi]) if hi is not None else (0, 0, 0)
                infos.append({"name": ws.title, "usable": hi is not None})
                if hi is not None and (best is None or sc > best[0]):
                    best = (sc, ws.title)
            if sheet is not None:
                if sheet not in wb.sheetnames:
                    raise GfaUploadError(f"'{sheet}' 시트가 없습니다", sheets=infos)
                name = sheet
            elif best:
                name = best[1]
            elif len(wb.worksheets) == 1:
                name = wb.worksheets[0].title
            else:
                raise GfaUploadError("성과 표가 있는 시트를 찾지 못했습니다 — 시트를 골라 주세요", sheets=infos)
            ws = wb[name]
            top = _ws_rows(ws, max_row=HEADER_SCAN)
            width = max((len(r) for r in top), default=0) or None
            rows = _ws_rows(ws, max_col=width)
        finally:
            wb.close()
        return Grid(kind="xlsx", rows=rows, sheet=name, sheets=infos)
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise GfaUploadError("옛 엑셀(.xls) 형식은 읽지 못합니다 — 엑셀에서 열어 .xlsx나 CSV로 다시 저장해 올려 주세요")
    if b"<table" in head or b"<html" in head or b"<?xml" in head and b"<table" in data[:65536].lower():
        rows, enc = _html_rows(data)
        return Grid(kind="html", rows=rows, encoding=enc)
    rows, enc, delim = _csv_rows(data)
    return Grid(kind="csv", rows=rows, encoding=enc, delimiter=delim)


# ─── 파일 → 일별 행 ─────────────────────────────────────────────────────────

@dataclass
class Record:
    account: str
    date: date
    campaign_id: Optional[str] = None
    campaign_name: Optional[str] = None
    adgroup_id: Optional[str] = None
    adgroup_name: Optional[str] = None
    creative_id: Optional[str] = None
    creative_name: Optional[str] = None
    creative_type: Optional[str] = None
    product: Optional[str] = None
    promotion: Optional[str] = None
    impressions: Optional[float] = None
    clicks: Optional[float] = None
    spend: float = 0.0
    conversions: Optional[float] = None
    conv_sales: Optional[float] = None
    budget: Optional[float] = None

    def as_dict(self) -> dict:
        d = {k: getattr(self, k) for k in ("account",) + DIMS + METRICS + ("budget",)}
        d["date"] = self.date.isoformat()
        return d


@dataclass
class Parsed:
    grid: Grid
    header_row: int  # 1부터 (파일 행 번호)
    mapping: Mapping
    records: list  # [Record] 같은 키는 합친 뒤
    dates: dict  # 계정 → {날짜} (값이 0인 날도 포함 — 다시 올릴 때 이 날짜들을 바꾼다)
    data_rows: int  # 읽은 데이터 행 수
    skipped: dict
    warnings: list
    errors: list
    spend_vat: Optional[str]
    budget_kind: Optional[str]  # daily | total
    budget_scope: Optional[str]  # campaign | adgroup

    @property
    def missing(self) -> list:
        return [f for f in REQUIRED if f not in self.mapping.fields]

    @property
    def can_apply(self) -> bool:
        return not self.missing and not self.errors and bool(self.records or self.dates)

    @property
    def granularity(self) -> str:
        f = self.mapping.fields
        if "creative_name" in f or "creative_id" in f:
            return "creative"
        if "adgroup_name" in f or "adgroup_id" in f:
            return "adgroup"
        if "campaign_name" in f or "campaign_id" in f:
            return "campaign"
        return "account"

    @property
    def date_range(self) -> tuple:
        ds = [d for s in self.dates.values() for d in s]
        return (min(ds), max(ds)) if ds else (None, None)

    def totals(self) -> dict:
        t = {m: 0.0 for m in METRICS}
        for r in self.records:
            for m in METRICS:
                t[m] += getattr(r, m) or 0.0
        out = {m: round(v, 2) for m, v in t.items()}
        for m in ("impressions", "clicks", "conversions", "conv_sales"):
            if m not in self.mapping.fields:
                out[m] = None
        sp = t["spend"]
        out["roas"] = round(t["conv_sales"] / sp, 4) if sp and out["conv_sales"] is not None else None
        return out


def _budget_meta(header: str) -> tuple:
    k = hkey(header)
    kind = "total" if re.search(r"총예산|기간예산|월예산|totalbudget|lifetimebudget", k) else "daily"
    scope = "adgroup" if re.search(r"광고그룹|adgroup|adset|광고세트", k) else "campaign"
    return kind, scope


def parse(data: bytes, filename: str = "", sheet: Optional[str] = None, header_row: Optional[int] = None,
          manual: Optional[dict] = None, account_label: str = "") -> Parsed:
    """파일 바이트 → 일별 행(합친 뒤)과 미리보기 재료. 오류는 errors에 모으고, 못 여는 파일만 예외."""
    grid = read_grid(data, filename, sheet)
    rows = grid.rows
    if not rows:
        raise GfaUploadError("파일에 내용이 없습니다")
    if header_row is not None:
        if not 1 <= header_row <= min(len(rows), 200):
            raise GfaUploadError(f"머리글 행 번호가 범위를 벗어났습니다 (1~{min(len(rows), 200)})")
        hi = header_row - 1
    else:
        hi = find_header(rows)
        if hi is None:
            hi = 0
    header = list(rows[hi])
    while header and (header[-1] is None or norm(header[-1]) == ""):
        header.pop()
    cols = [classify(i, v) for i, v in enumerate(header)]
    mp = pick_columns(cols, manual)
    f = mp.fields
    warnings: list = list(mp.notes)
    errors: list = []

    spend_vat = vat_of(cols[f["spend"][0]].header) if "spend" in f else None
    budget_kind = budget_scope = None
    if "budget" in f:
        budget_kind, budget_scope = _budget_meta(cols[f["budget"][0]].header)

    label = norm(account_label)[:200]
    acc_col = f.get("account", [None])[0] if "account" in f else None
    agg: dict = {}
    dates: dict = {}
    skipped = {"blank": 0, "total": 0, "no_date": 0, "bad_date": 0, "range": 0, "zero": 0}
    bad_examples: list = []
    range_examples: list = []
    nonnum: dict = {}
    data_rows = 0

    if "date" in f and "spend" in f:
        di = f["date"][0]
        for ri in range(hi + 1, len(rows)):
            row = rows[ri]
            if not any(v is not None and norm(v) != "" for v in row):
                skipped["blank"] += 1
                continue
            data_rows += 1
            cell = lambda i: row[i] if i < len(row) else None  # noqa: E731
            st, d = parse_day(cell(di))
            if st == "total":
                skipped["total"] += 1
                continue
            if st == "empty":
                skipped["no_date"] += 1  # 맨 아래 합계 행 등
                continue
            if st == "range":
                skipped["range"] += 1
                if len(range_examples) < 3:
                    range_examples.append(f"{ri + 1}행 「{norm(cell(di))}」")
                continue
            if st == "bad":
                skipped["bad_date"] += 1
                if len(bad_examples) < 3:
                    bad_examples.append(f"{ri + 1}행 「{norm(cell(di))[:40]}」")
                continue
            dims = {}
            total_row = False
            for k in DIMS:
                if k in f:
                    v = cell(f[k][0])
                    if k in ("campaign_name", "adgroup_name", "creative_name") and is_total(v):
                        total_row = True
                    dims[k] = _text(v, 100 if k.endswith("_id") or k == "creative_type" else 300)
            # 캠페인 칸이 '전체'이고 아래 단위 칸이 비면 날짜별 합계 행
            if (dims.get("campaign_name") == "전체"
                    and not any(dims.get(k) for k in ("adgroup_id", "adgroup_name", "creative_id", "creative_name"))):
                total_row = True
            if total_row:
                skipped["total"] += 1
                continue
            acc = label or (_text(cell(acc_col), 200) if acc_col is not None else None) or ""
            dates.setdefault(acc, set()).add(d)
            vals: dict = {}
            for k in METRICS + ("budget",):
                if k not in f:
                    continue
                s = 0.0
                for i in f[k]:
                    x, ok = parse_num(cell(i))
                    if not ok:
                        h = cols[i].header or col_letter(i)
                        nonnum[h] = nonnum.get(h, 0) + 1
                    s += x
                vals[k] = s
            if not any(vals.get(k) for k in METRICS) and not vals.get("budget"):
                skipped["zero"] += 1
                continue
            key = (acc, d) + tuple(dims.get(k) for k in DIMS)
            rec = agg.get(key)
            if rec is None:
                rec = Record(account=acc, date=d, **dims)
                for k in ("impressions", "clicks", "conversions", "conv_sales"):
                    if k in f:
                        setattr(rec, k, 0.0)
                agg[key] = rec
            for k in METRICS:
                if k in vals:
                    setattr(rec, k, (getattr(rec, k) or 0.0) + vals[k])
            if vals.get("budget"):
                rec.budget = max(rec.budget or 0.0, vals["budget"])
        if data_rows > MAX_DATA_ROWS:
            errors.append(f"행이 너무 많습니다 (최대 {MAX_DATA_ROWS:,}행) — 기간을 나눠 올려 주세요")

    missing = [k for k in REQUIRED if k not in f]
    if missing:
        errors.append("필수 칸을 찾지 못했습니다: " + ", ".join(FIELDS[k] for k in missing)
                      + " — 아래에서 칸을 직접 골라 주세요")
    if range_examples:
        errors.append(f"날짜 칸에 여러 날을 묶은 기간이 {skipped['range']}행 있습니다(" + ", ".join(range_examples)
                      + ") — 보고서 기간 단위를 '일'로 두고 내려받아 주세요")
    if not missing and not range_examples and not agg and not dates:
        errors.append("읽을 수 있는 일별 행이 없습니다 — 날짜·비용 칸이 맞는지 확인해 주세요")
    if bad_examples:
        warnings.append(f"날짜를 읽지 못한 행 {skipped['bad_date']}개는 건너뜁니다 (예: "
                        + ", ".join(bad_examples) + ")")
    if skipped["total"]:
        warnings.append(f"합계·소계 행 {skipped['total']}개는 건너뜁니다")
    if skipped["no_date"]:
        warnings.append(f"날짜가 빈 행 {skipped['no_date']}개는 건너뜁니다 (보통 맨 아래 합계 행)")
    if nonnum:
        ex = ", ".join(f"「{h}」 {n}칸" for h, n in list(nonnum.items())[:3])
        warnings.append(f"숫자로 읽지 못한 칸은 0으로 봤습니다: {ex}")
    if "spend" in f:
        if spend_vat == "incl":
            warnings.append(f"비용 「{cols[f['spend'][0]].header}」은 VAT 포함 값입니다 (VAT 제외 비용 칸이 없음)")
        elif spend_vat is None:
            warnings.append(f"비용 「{cols[f['spend'][0]].header}」 머리글에 VAT 표시가 없어 파일 값 그대로 씁니다")
    if "date" in f and "spend" in f:
        if "conv_sales" not in f:
            warnings.append("전환 매출 칸이 없습니다 — ROAS·전환 매출이 비게 됩니다 (칸을 직접 고르거나 전환 매출이 든 보고서로 받아 주세요)")
        if "conversions" not in f:
            warnings.append("전환 수 칸이 없습니다 — CPA·전환율이 비게 됩니다")
        if "impressions" not in f or "clicks" not in f:
            warnings.append("노출·클릭 칸이 없어 CTR·CPC 일부가 비게 됩니다")
    br = [c.header for c in cols if c.breakdown]
    if br:
        warnings.append(f"「{', '.join(br)}」로 나뉜 행은 합쳐서 봅니다")
    if budget_kind == "total":
        warnings.append("예산 칸이 총(기간) 예산입니다 — 예산 소진율은 선택 기간 광고비 ÷ 총 예산으로 계산합니다")

    records = sorted(agg.values(), key=lambda r: (r.date, r.account, -(r.spend or 0)))
    return Parsed(grid=grid, header_row=hi + 1, mapping=mp, records=records, dates=dates,
                  data_rows=data_rows, skipped=skipped, warnings=warnings, errors=errors,
                  spend_vat=spend_vat, budget_kind=budget_kind, budget_scope=budget_scope)


def columns_info(p: Parsed) -> list:
    """미리보기 표 — 칸마다 (열, 머리글, 쓰는 항목)."""
    used = {i: k for k, idxs in p.mapping.fields.items() for i in idxs}
    out = []
    for c in p.mapping.cols:
        k = used.get(c.index)
        out.append({
            "index": c.index, "letter": col_letter(c.index), "header": c.header,
            "field": k, "field_label": FIELDS.get(k) if k else None,
            "known": c.field is not None,  # 알아본 칸(쓰지 않아도)
            "status": "used" if k else ("ignored" if c.field else ("empty" if not c.header else "unknown")),
        })
    return out


def mapping_meta(p: Parsed) -> dict:
    """반영 기록에 남길 칸 연결 요약."""
    cols = p.mapping.cols
    return {
        "fields": {k: [{"index": i, "header": cols[i].header} for i in idxs] for k, idxs in p.mapping.fields.items()},
        "manual": sorted(p.mapping.manual),
        "header_row": p.header_row,
        "kind": p.grid.kind,
        "spend_vat": p.spend_vat,
        "budget_kind": p.budget_kind,
        "budget_scope": p.budget_scope,
    }
